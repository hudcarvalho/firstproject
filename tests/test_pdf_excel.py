from decimal import Decimal as D
from io import BytesIO

from openpyxl import load_workbook

from app.pdf_excel import numero_br
from app.pdf_simples import Documento


def _pdf_relatorio():
    """Relatório sem bordas: título, cabeçalho e linhas com código, descrição e dois valores."""
    doc = Documento("Balancete")
    doc.texto(40, 40, "Balancete de verificação — Empresa Teste", 13, negrito=True)
    y = 80
    for x, t, al in ((40, "Conta", "esquerda"), (100, "Descrição", "esquerda"),
                     (420, "Débito", "direita"), (540, "Crédito", "direita")):
        doc.texto(x, y, t, 9, negrito=True, alinhar=al)
    linhas = [("1.1.01", "Caixa e equivalentes", "12.345,67", ""),
              ("1.1.02", "Bancos conta movimento", "1.000,00", "(250,50)"),
              ("2.1.01", "Fornecedores", "", "98.765,43"),
              ("3.1.01", "Receita de vendas — ação", "0,00", "1.234.567,89")]
    for conta, desc, deb, cred in linhas:
        y += 14
        doc.texto(40, y, conta, 9)
        doc.texto(100, y, desc, 9)
        doc.texto(420, y, deb, 9, alinhar="direita")
        doc.texto(540, y, cred, 9, alinhar="direita")
    return doc.gerar()


def test_numero_br():
    assert numero_br("1.234,56") == D("1234.56")
    assert numero_br("(1.234,56)") == D("-1234.56")
    assert numero_br("-1.234,56") == D("-1234.56") and numero_br("1.234,56-") == D("-1234.56")
    assert numero_br("R$ 10,00") == D("10.00") and numero_br("1.234,56 D") == D("1234.56")
    for texto in ("12.345.678/0001-90", "01/2026", "1001", "abc", ""):
        assert numero_br(texto) is None


def test_tela_e_menu(client):
    assert "Converter PDF em Excel" in client.get("/empresas").get_data(as_text=True)
    assert 'enctype="multipart/form-data"' in client.get("/ferramentas/pdf-excel").get_data(as_text=True)


def test_converte_relatorio_em_colunas(client):
    r = client.post("/ferramentas/pdf-excel", data={
        "arquivo": (BytesIO(_pdf_relatorio()), "Balancete março.pdf"), "modo": "layout", "numeros": "on",
    }, content_type="multipart/form-data")
    assert r.status_code == 200 and "Balancete%20mar%C3%A7o.xlsx" in r.headers["Content-Disposition"]
    assert r.headers["X-Paginas"] == "1"
    linhas = list(load_workbook(BytesIO(r.data)).active.iter_rows(values_only=True))
    assert linhas[0][0] == "Balancete de verificação — Empresa Teste"
    assert linhas[1] == ("Conta", "Descrição", "Débito", "Crédito")
    assert linhas[2] == ("1.1.01", "Caixa e equivalentes", 12345.67, None)
    assert linhas[3] == ("1.1.02", "Bancos conta movimento", 1000, -250.5)
    assert linhas[5] == ("3.1.01", "Receita de vendas — ação", 0, 1234567.89)


def test_sem_converter_numeros_e_aba_por_pagina(client):
    r = client.post("/ferramentas/pdf-excel", data={
        "arquivo": (BytesIO(_pdf_relatorio()), "x.pdf"), "modo": "layout", "aba_por_pagina": "on",
    }, content_type="multipart/form-data")
    wb = load_workbook(BytesIO(r.data))
    assert wb.sheetnames == ["Página 1"]
    assert list(wb.active.iter_rows(values_only=True))[2][2] == "12.345,67"


def test_erros(client):
    r = client.post("/ferramentas/pdf-excel", data={"arquivo": (BytesIO(b"nao sou pdf"), "a.pdf")},
                    content_type="multipart/form-data", follow_redirects=True)
    assert "não é um PDF" in r.get_data(as_text=True)
    r = client.post("/ferramentas/pdf-excel", data={"arquivo": (BytesIO(b"%PDF-1.4 quebrado"), "a.pdf")},
                    content_type="multipart/form-data", follow_redirects=True)
    assert "Não foi possível ler" in r.get_data(as_text=True)
    # relatório sem bordas no modo "somente tabelas": nada a extrair
    r = client.post("/ferramentas/pdf-excel", data={"arquivo": (BytesIO(_pdf_relatorio()), "a.pdf"), "modo": "tabelas"},
                    content_type="multipart/form-data", follow_redirects=True)
    assert "Nenhum texto encontrado" in r.get_data(as_text=True)


def _pdf_extrato_itau():
    """Extrato mensal no layout do Itaú (posições das colunas iguais às do banco), dados fictícios."""
    doc = Documento("Extrato")
    t = doc.texto
    t(40, 30, "extrato mensal ag 1234 cc 12345-6 jul 2026 001|002", 8)
    t(40, 50, "CTC CENTRO", 8)
    t(40, 62, "EMPRESA FICTICIA LTDA", 8)
    t(40, 74, "AV PAULISTA 1000", 8)
    t(411, 110, "saldo em 30/06/26 saldo em 31/07/26", 8)
    t(58, 122, "01. Conta Corrente e Aplicações Automáticas", 9)
    t(421, 134, "R$ 1.000,00 R$ 1.500,00", 9)
    t(343, 150, "entradas (créditos)", 8)
    t(343, 162, "total 2.000,00", 8)
    t(343, 174, "saídas (débitos)", 8)
    t(343, 186, "total 1.500,00", 8)
    t(56, 210, "Conta Corrente | Movimentação", 9)
    for x, texto in ((150, "data"), (208, "descrição"), (357, "entradas"), (388, "R$"), (426, "saídas"),
                     (448, "R$"), (521, "saldo"), (541, "R$")):
        t(x, 230, texto, 8)
    linhas = [
        ("30/06", "Saldo anterior", None, None, "1.000,00"),
        ("01/07", "Sispag CLIENTE XPTO", "2.000,00", None, None),
        ("", "BOLETO PAGO FORNECEDOR", None, "500,00-", None),
        ("", "Apl Aplic Aut Mais", None, "1.499,00-", "1,00"),
        ("", "SALDO APLIC AUT MAIS", None, None, "2.499,00"),
        ("02/07", "PIX ENVIADO JOAO SILVA", None, "1.000,00-", None),
        ("", "Res Aplic Aut Mais", "1.000,00", None, "1,00"),
        ("", "SALDO APLIC AUT MAIS", None, None, "1.499,00"),
        ("31/07", "Saldo em C/C", None, None, "1,00"),
        ("", "Saldo final", None, None, "1.500,00"),
    ]
    y = 250
    for data, desc, ent, sai, saldo in linhas:
        y += 12
        t(150, y, data, 8)
        t(208, y, desc, 8)
        t(396, y, ent, 8, alinhar="direita")
        t(458, y, sai, 8, alinhar="direita")
        t(549, y, saldo, 8, alinhar="direita")
    doc.nova_pagina()
    t(40, 40, "CDB e Renda Fixa", 9)
    t(40, 60, "CDB-DI 15/02/22 15/02/27 10.000,00 100,00% DI 12.000,00 11.500,00", 8)
    t(40, 72, "total aplicações financeiras 10.000,00 12.000,00 11.500,00", 8)
    return doc.gerar()


def test_extrato_itau_organizado(client):
    r = client.post("/ferramentas/pdf-excel", data={"arquivo": (BytesIO(_pdf_extrato_itau()), "extrato.pdf")},
                    content_type="multipart/form-data")
    wb = load_workbook(BytesIO(r.data))
    assert wb.sheetnames == ["Resumo", "Extrato", "Investimentos CDB"]
    ext = wb["Extrato"]
    assert ext["A1"].value == "EMPRESA FICTICIA LTDA — Itaú ag. 1234 c/c 12345-6"
    linhas = [tuple(c.value for c in ext[r][:9]) for r in range(6, 9)]
    # as linhas internas da aplicação automática ficam de fora
    assert [l[1:6] for l in linhas] == [
        ("Sispag CLIENTE XPTO", "CLIENTE XPTO", "Recebimento Sispag", 2000, None),
        ("BOLETO PAGO FORNECEDOR", "FORNECEDOR", "Boletos/contas pagas", None, 500),
        ("PIX ENVIADO JOAO SILVA", "JOAO SILVA", "PIX enviado", None, 1000),
    ]
    # saldo do banco no fim do dia = aplicação automática + conta corrente, com conferência
    assert linhas[1][7] == 2500 and linhas[2][7] == 1500 and linhas[0][7] is None
    assert linhas[2][6] == "=G7+E8-F8" and linhas[2][8].startswith("=IF(ROUND(G8-H8,2)=0")
    assert ext["B9"].value == "TOTAIS DO MÊS" and ext["E9"].value == "=SUM(E6:E8)"
    res = wb["Resumo"]
    assert [res.cell(i, 3).value for i in range(4, 8)] == [1000, 2000, 1500, 1500]
    assert res["A13"].value == "Recebimento Sispag" and res["B13"].value.startswith("=SUMIFS(")
    inv = wb["Investimentos CDB"]
    assert [c.value for c in inv[4]][:7][3:] == [10000, 1, 12000, 11500]


def test_extrato_itau_no_modo_layout_fica_simples(client):
    r = client.post("/ferramentas/pdf-excel", data={"arquivo": (BytesIO(_pdf_extrato_itau()), "e.pdf"), "modo": "layout"},
                    content_type="multipart/form-data")
    assert load_workbook(BytesIO(r.data)).sheetnames == ["PDF"]
