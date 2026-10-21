import re
from decimal import Decimal as D

import pytest

from app.lalur import demonstrativos
from app.models import Empresa, LalurApuracao, RegimeTributario, db


@pytest.fixture
def empresa_lr(app):
    with app.app_context():
        lr = RegimeTributario.query.filter_by(nome="Lucro Real").one()
        e = Empresa(codigo="2120", razao_social="EBM COMERCIO DE CHOCOLATES LTDA", cnpj="11222333000181", regime=lr)
        db.session.add(e)
        db.session.commit()
        return e.id


def test_aba_so_no_lucro_real(client, app, empresa_lr):
    assert ">LALUR<" in client.get(f"/empresas/{empresa_lr}/contabil").get_data(as_text=True)
    with app.app_context():
        e = db.session.get(Empresa, empresa_lr)
        e.regime = RegimeTributario.query.filter_by(nome="Lucro Presumido").one()
        db.session.commit()
    assert ">LALUR<" not in client.get(f"/empresas/{empresa_lr}/contabil").get_data(as_text=True)
    assert client.get(f"/lalur/{empresa_lr}").status_code == 302


def test_apuracao_trimestral_encadeada(client, app, empresa_lr):
    url = f"/lalur/{empresa_lr}"
    client.post(f"{url}/saldo-inicial?ano=2026", data={"prejuizo_fiscal": "50.000,00", "base_negativa": "20.000"})
    t1 = f"{url}/2026/1"
    client.post(t1, data={"lucro_antes": "100.000,00", "compensar": "on", "aliquota_csll": "9",
                          "irpj_deduzir": "1.000,00"})
    client.post(f"{t1}/lancamentos", data={"tipo": "adicao", "tributo": "ambos", "descricao": "Multas", "valor": "10.000"})
    client.post(f"{t1}/lancamentos", data={"tipo": "exclusao", "tributo": "irpj", "descricao": "Equivalência", "valor": "5.000"})
    client.post(f"{url}/2026/2", data={"lucro_antes": "-30.000,00", "compensar": "on"})
    client.post(f"{url}/2026/3", data={"lucro_antes": "20.000,00", "compensar": "on"})

    with app.app_context():
        calc, (prejuizo, base_neg) = demonstrativos(db.session.get(Empresa, empresa_lr))
    d1 = calc[(2026, 1)]
    assert d1.irpj.ajustado == D("105000") and d1.irpj.compensacao == D("31500.00")
    assert d1.irpj.base == D("73500.00") and d1.irpj_normal == D("11025.00")
    assert d1.irpj_adicional == D("1350.00") and d1.irpj_a_pagar == D("11375.00")
    assert d1.csll.ajustado == D("110000") and d1.csll.compensacao == D("20000")
    assert d1.csll_devida == D("8100.00") and d1.csll.saldo_final == 0
    assert d1.lucro_liquido == D("100000") - D("12375.00") - D("8100.00")  # 79.525,00
    d2 = calc[(2026, 2)]
    assert d2.irpj.gerado == D("30000") and d2.irpj.saldo_final == D("48500.00")
    assert d2.irpj_a_pagar == 0 and d2.csll.saldo_final == D("30000")
    d3 = calc[(2026, 3)]
    assert d3.irpj.compensacao == D("6000.00") and d3.irpj_normal == D("2100.00") and d3.irpj_adicional == 0
    assert prejuizo == D("42500.00") and base_neg == D("24000.00")

    resumo = client.get(f"{url}?ano=2026").get_data(as_text=True)
    assert "11.375,00" in resumo and "42.500,00" in resumo
    assert "= Lucro líquido do período" in resumo and "79.525,00" in resumo
    tela = client.get(t1).get_data(as_text=True)
    assert "Multas" in tela and "73.500,00" in tela


def test_nao_mistura_anual_e_trimestral(client, app, empresa_lr):
    url = f"/lalur/{empresa_lr}"
    client.post(f"{url}/2026/1", data={"lucro_antes": "1.000"})
    r = client.post(f"{url}/2026/0", data={"lucro_antes": "1.000"}, follow_redirects=True)
    assert "outro formato" in r.get_data(as_text=True)
    # anual: adicional sobre o que passar de 240 mil
    client.post(f"{url}/2025/0", data={"lucro_antes": "300.000,00"})
    with app.app_context():
        calc, _ = demonstrativos(db.session.get(Empresa, empresa_lr))
        assert calc[(2025, 0)].irpj_adicional == D("6000.00")
        assert LalurApuracao.query.filter_by(ano=2026, periodo=0).count() == 0


def test_lancamento_invalido_e_remocao(client, app, empresa_lr):
    t1 = f"/lalur/{empresa_lr}/2026/1"
    r = client.post(f"{t1}/lancamentos", data={"tipo": "adicao", "descricao": "X", "valor": "-5"}, follow_redirects=True)
    assert "Lançamento inválido" in r.get_data(as_text=True)
    client.post(f"{t1}/lancamentos", data={"tipo": "adicao", "descricao": "Brindes", "valor": "500"})
    with app.app_context():
        lid = LalurApuracao.query.one().lancamentos[0].id
    client.post(f"/lalur/lancamentos/{lid}/remover")
    with app.app_context():
        assert LalurApuracao.query.one().lancamentos == []
    client.post(f"{t1}/excluir")
    with app.app_context():
        assert LalurApuracao.query.count() == 0


def test_formularios_lalur_com_csrf(client, app, empresa_lr):
    client.post(f"/lalur/{empresa_lr}/2026/1/lancamentos", data={"tipo": "adicao", "descricao": "Brindes", "valor": "500"})
    app.config["WTF_CSRF_ENABLED"] = True
    for url in [f"/lalur/{empresa_lr}?ano=2026", f"/lalur/{empresa_lr}/2026/1"]:
        html = client.get(url).get_data(as_text=True)
        for form in re.findall(r"<form\b[^>]*>.*?</form>", html, flags=re.S | re.I):
            if re.search(r'method="post"', form, flags=re.I):
                assert 'name="csrf_token"' in form, form[:160]


def test_exportar_pdf_e_excel(client, app, empresa_lr):
    from io import BytesIO

    from openpyxl import load_workbook

    url = f"/lalur/{empresa_lr}"
    assert client.get(f"{url}/2026/exportar").status_code == 302  # nada apurado ainda
    client.post(f"{url}/2026/1", data={"lucro_antes": "29.480,47"})
    client.post(f"{url}/2026/1/lancamentos", data={"tipo": "adicao", "descricao": "Multas e ação fiscal", "valor": "1.000"})
    client.post(f"{url}/2026/2", data={"lucro_antes": "70.455,22"})
    client.post(f"{url}/2026/3/lancamentos", data={"tipo": "adicao", "descricao": "Sem lucro", "valor": "1"})

    tela = client.get(f"{url}?ano=2026").get_data(as_text=True)
    assert "Todos os períodos apurados de 2026 (2)" in tela and "Só o 2º trim." in tela
    assert "Só o 3º trim." not in tela  # sem lucro informado não conta como apurado

    r = client.get(f"{url}/2026/exportar?periodo=todos&formato=pdf")
    assert r.mimetype == "application/pdf" and r.data.startswith(b"%PDF") and r.data.rstrip().endswith(b"%%EOF")
    assert "lalur_2120_2026_todos.pdf" in r.headers["Content-Disposition"]
    assert "1º trim.".encode("cp1252") in r.data and "2º trim.".encode("cp1252") in r.data
    assert "Multas e ação fiscal".encode("cp1252") in r.data

    r = client.get(f"{url}/2026/exportar?periodo=2&formato=xlsx")
    assert "lalur_2120_2026_2tri.xlsx" in r.headers["Content-Disposition"]
    wb = load_workbook(BytesIO(r.data))
    ws = wb["Parte A"]
    linhas = list(ws.iter_rows(values_only=True))
    assert linhas[4] == ("R$", "2º trim.", "Fórmula / critério")
    por_rotulo = {l[0]: l[1:] for l in linhas if l[0]}
    assert por_rotulo["Lucro líquido antes do IRPJ/CSLL"][0] == 70455.22
    # o cálculo é feito por fórmulas do Excel, com a explicação ao lado
    assert por_rotulo["= Lucro ajustado"][0].startswith("=B")
    formula_irpj, explicacao = por_rotulo["IRPJ 15%"]
    assert formula_irpj.startswith("=ROUND(") and "'Parâmetros'!$B$2" in formula_irpj
    assert explicacao == "Lucro real × 15%"
    assert por_rotulo["= Lucro líquido do período"][0].startswith("=")
    assert por_rotulo["(+) Adições"][0].startswith("=SUMIFS(")
    assert wb["Parâmetros"]["B2"].value == 0.15 and wb["Parâmetros"]["B5"].value == 0.3

    r = client.get(f"{url}/2026/exportar?periodo=todos&formato=xlsx")
    wb = load_workbook(BytesIO(r.data))
    ws = wb["Parte A"]
    assert list(ws.iter_rows(values_only=True))[4][:3] == ("R$", "1º trim.", "2º trim.")
    saldo_anterior = next(r for r in ws.iter_rows() if r[0].value == "Prejuízo fiscal a compensar (saldo anterior)")
    assert saldo_anterior[1].value == 0 and saldo_anterior[2].value.startswith("=B")  # encadeado
    lanc = list(wb["Adições e exclusões"].iter_rows(values_only=True))
    assert lanc[5] == ("1º trim.", "Adição", "Multas e ação fiscal", "IRPJ e CSLL", 1000)


def test_guia_darf_quotas_e_pdf(client, app, empresa_lr):
    from datetime import date

    from app.guia_darf import quotas, ultimo_dia_util

    assert ultimo_dia_util(2026, 5) == date(2026, 5, 29)  # 31/05/2026 é domingo
    lista, aviso = quotas(D("1500"), 2026, 2, 3)
    assert len(lista) == 1 and "quota única" in aviso
    lista, _ = quotas(D("7793.28"), 2026, 2, 3, selic=D("1.05"))
    assert [q.vencimento for q in lista] == [date(2026, 7, 31), date(2026, 8, 31), date(2026, 9, 30)]
    assert sum(q.principal for q in lista) == D("7793.28")
    assert lista[1].juros == D("25.98") and lista[2].juros_pct == D("2.05")
    assert quotas(D("5000"), 2026, 0, 3)[0][0].vencimento == date(2027, 3, 31)

    url = f"/lalur/{empresa_lr}/2026/2"
    assert client.get(f"{url}/guia").status_code == 302  # nada apurado
    client.post(url, data={"lucro_antes": "70.455,22"})
    assert "Guia (DARF)" in client.get(url).get_data(as_text=True)
    tela = client.get(f"{url}/guia?forma=parcelado&parcelas=3").get_data(as_text=True)
    assert "0220" in tela and "6012" in tela and "3ª de 3" in tela and "Selic + 1%" in tela
    assert "código de barras" in tela
    r = client.get(f"{url}/guia?forma=unica&baixar=1")
    assert r.mimetype == "application/pdf" and r.data.startswith(b"%PDF")
    assert "darf_dados_2120_2026_2tri_unica.pdf" in r.headers["Content-Disposition"]
    assert b"11.613,80" in r.data and "NÃO é documento de arrecadação".encode("cp1252") in r.data
