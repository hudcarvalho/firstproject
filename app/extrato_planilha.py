"""Planilha organizada de extrato bancário (comum aos bancos reconhecidos).

Abas:
- Resumo: saldos, entradas e saídas calculados x impressos pelo banco (diferença) e a
  movimentação por tipo (SOMASES / CONT.SE);
- Extrato: data, histórico, [documento], favorecido, tipo, entrada, saída, saldo por fórmula e
  o saldo impresso pelo banco com "Confere";
- abas extras conforme o banco (Investimentos CDB, Invest Fácil, Últimos lançamentos).
"""
import unicodedata
from dataclasses import dataclass, field
from datetime import date, timedelta
from decimal import Decimal
from io import BytesIO

NOMES_MES = ["janeiro", "fevereiro", "março", "abril", "maio", "junho", "julho", "agosto",
             "setembro", "outubro", "novembro", "dezembro"]

# Tipos na ordem em que aparecem no resumo.
TIPOS = [
    "Aplicação em CDB", "Resgate de CDB", "Recebimento PIX/TED", "Recebimento Sispag",
    "Cobrança (títulos liquidados)", "Depósitos", "Transferência recebida", "Consórcio",
    "Tributos/guias", "Transferência CC Itaú", "Transferência entre contas", "PIX enviado",
    "TED/DOC enviado", "Pagamento Sispag", "Débito automático", "Boletos/contas pagas",
    "Cartão de crédito", "Tarifas bancárias", "Juros/IOF", "Rendimento aplicação automática",
    "Rendimento Invest Fácil", "Outras entradas", "Outras saídas",
]


def sem_acento(t):
    return "".join(c for c in unicodedata.normalize("NFD", t) if unicodedata.category(c) != "Mn").upper()


def valor(t):
    """'1.234,56', '-1.234,56' ou '1.234,56-' -> Decimal."""
    neg = t.startswith("-") or t.endswith("-")
    v = Decimal(t.strip("-").replace(".", "").replace(",", "."))
    return -v if neg else v


@dataclass
class Movimento:
    data: date
    historico: str
    entrada: Decimal | None = None
    saida: Decimal | None = None
    saldo_banco: Decimal | None = None  # saldo impresso pelo banco (fim do dia ou da linha)
    favorecido: str = ""
    tipo: str = ""
    documento: str = ""


@dataclass
class Extrato:
    banco: str = ""
    empresa: str = ""
    agencia: str = ""
    conta: str = ""
    inicio: date | None = None
    fim: date | None = None
    data_saldo_anterior: date | None = None
    saldo_anterior: Decimal | None = None
    saldo_final: Decimal | None = None
    total_entradas: Decimal | None = None
    total_saidas: Decimal | None = None
    tem_aplicacao_automatica: bool = False
    nota_saldo: str = ""  # complemento do subtítulo do Extrato
    origem_resumo: str = "no extrato"  # onde o banco imprime os totais (nota do Resumo)
    rotulo_favorecido: str = "Favorecido / Remetente"
    rotulo_saldo_banco: str = "Saldo banco fim do dia (R$)"
    movimentos: list = field(default_factory=list)
    cdbs: list = field(default_factory=list)  # dicts (Itaú)
    total_cdb: tuple | None = None
    saldos_invest: list = field(default_factory=list)  # [(data, histórico, valor)] (Bradesco)
    ultimos: list = field(default_factory=list)  # lançamentos posteriores ao período
    ultimos_saldo_anterior: Decimal | None = None
    ultimos_data_anterior: date | None = None


def _br(v):
    return f"{v:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def planilha(ex):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.workbook.properties import CalcProperties

    MOEDA = '#,##0.00;[Red]\\-#,##0.00;\\-'
    cab_fonte = Font(bold=True, color="FFFFFF", size=11)
    cab_fundo = PatternFill("solid", fgColor="1F3864")
    azul = Font(color="0000FF", size=10)
    normal = Font(size=10)
    negrito = Font(bold=True, size=10)
    cinza = PatternFill("solid", fgColor="F2F2F2")
    fim = ex.fim or (ex.movimentos[-1].data if ex.movimentos else None)
    inicio = ex.inicio
    mes_ano = f"{NOMES_MES[fim.month - 1]}/{fim.year}" if fim else ""
    conta = f"{ex.banco} ag. {ex.agencia} c/c {ex.conta}"
    titulo = " — ".join(x for x in (ex.empresa, conta) if x)
    com_documento = any(m.documento for m in ex.movimentos + ex.ultimos)

    # colunas do Extrato
    nomes = ["Data", "Histórico (banco)"] + (["Dcto."] if com_documento else []) + [
        ex.rotulo_favorecido, "Tipo", "Entrada (R$)", "Saída (R$)", "Saldo (R$)", ex.rotulo_saldo_banco, "Confere"]
    C = {chave: i + 1 for i, chave in enumerate(
        ["data", "hist"] + (["doc"] if com_documento else []) +
        ["fav", "tipo", "ent", "sai", "saldo", "banco", "confere"])}
    L = {k: get_column_letter(v) for k, v in C.items()}
    larguras = {"data": 11, "hist": 32, "doc": 11, "fav": 30, "tipo": 30, "ent": 15, "sai": 15, "saldo": 15,
                "banco": 16, "confere": 9}

    def cabecalho(ws, linha, textos):
        for i, t in enumerate(textos, start=1):
            c = ws.cell(linha, i, t)
            c.font, c.fill = cab_fonte, cab_fundo
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    def fmt(c, formato=MOEDA, fonte=normal):
        c.number_format = formato
        c.font = fonte
        return c

    def aba_extrato(ws, subtitulo, movimentos, saldo_anterior, data_anterior):
        """Escreve um extrato (linha 5 = saldo anterior) e devolve (última linha, linha de totais)."""
        ws["A1"] = titulo
        ws["A1"].font = Font(bold=True, size=12)
        ws["A2"] = subtitulo
        ws["A2"].font = normal
        cabecalho(ws, 4, nomes)
        ws.cell(5, C["data"], data_anterior).number_format = "dd/mm/yyyy"
        ws.cell(5, C["hist"], "SALDO ANTERIOR")
        ws.cell(5, C["saldo"], saldo_anterior)
        for col in range(1, len(nomes) + 1):
            ws.cell(5, col).font, ws.cell(5, col).fill = negrito, cinza
        ws.cell(5, C["saldo"]).number_format = MOEDA
        ws.cell(5, C["saldo"]).font = Font(bold=True, size=10, color="0000FF")
        r = 5
        for m in movimentos:
            r += 1
            ws.cell(r, C["data"], m.data).number_format = "dd/mm/yyyy"
            ws.cell(r, C["data"]).font = normal
            ws.cell(r, C["hist"], m.historico).font = normal
            if com_documento:
                ws.cell(r, C["doc"], m.documento or None).font = normal
            ws.cell(r, C["fav"], m.favorecido or None).font = normal
            ws.cell(r, C["tipo"], m.tipo).font = normal
            fmt(ws.cell(r, C["ent"], m.entrada))
            fmt(ws.cell(r, C["sai"], m.saida))
            fmt(ws.cell(r, C["saldo"], f"={L['saldo']}{r - 1}+{L['ent']}{r}-{L['sai']}{r}"))
            if m.saldo_banco is not None:
                fmt(ws.cell(r, C["banco"], m.saldo_banco), fonte=azul)
                ws.cell(r, C["confere"], f'=IF(ROUND({L["saldo"]}{r}-{L["banco"]}{r},2)=0,"OK","DIVERGE")').font = normal
        ultima = r
        total = ultima + 1
        ws.cell(total, C["hist"], "TOTAIS DO MÊS" if movimentos is ex.movimentos else "TOTAIS")
        fmt(ws.cell(total, C["ent"], f"=SUM({L['ent']}6:{L['ent']}{ultima})"))
        fmt(ws.cell(total, C["sai"], f"=SUM({L['sai']}6:{L['sai']}{ultima})"))
        fmt(ws.cell(total, C["saldo"], f"={L['saldo']}{ultima}"))
        for col in range(1, len(nomes) + 1):
            ws.cell(total, col).font = negrito
            ws.cell(total, col).fill = cinza
        for chave, col in C.items():
            ws.column_dimensions[get_column_letter(col)].width = larguras[chave]
        ws.row_dimensions[4].height = 30
        ws.freeze_panes = "A5"
        ws.auto_filter.ref = f"A4:{get_column_letter(len(nomes))}{ultima}"
        return ultima, total

    wb = Workbook()
    resumo = wb.active
    resumo.title = "Resumo"
    ext = wb.create_sheet("Extrato")

    data_anterior = ex.data_saldo_anterior or (inicio - timedelta(days=1) if inicio else None)
    subtitulo = f"Extrato de {inicio:%d/%m/%Y} a {fim:%d/%m/%Y}" if inicio and fim else "Extrato"
    if ex.nota_saldo:
        subtitulo += f" ({ex.nota_saldo})"
    ultima, total = aba_extrato(ext, subtitulo, ex.movimentos, ex.saldo_anterior, data_anterior)

    # ---- Resumo
    resumo["A1"] = f"{titulo} — {mes_ano}" if mes_ano else titulo
    resumo["A1"].font = Font(bold=True, size=12)
    cabecalho(resumo, 3, ["Item", "Calculado (R$)", "Banco (R$)", "Diferença"])
    itens = [
        (f"Saldo em {data_anterior:%d/%m/%y}" if data_anterior else "Saldo anterior",
         f"=Extrato!{L['saldo']}5", ex.saldo_anterior),
        ("Entradas (créditos)", f"=Extrato!{L['ent']}{total}", ex.total_entradas),
        ("Saídas (débitos)", f"=Extrato!{L['sai']}{total}", ex.total_saidas),
        (f"Saldo em {fim:%d/%m/%y}" if fim else "Saldo final", f"=Extrato!{L['saldo']}{ultima}", ex.saldo_final),
    ]
    for i, (rotulo, formula, banco) in enumerate(itens, start=4):
        resumo.cell(i, 1, rotulo).font = normal
        fmt(resumo.cell(i, 2, formula))
        fmt(resumo.cell(i, 3, banco), fonte=azul)
        fmt(resumo.cell(i, 4, f"=ROUND(B{i}-C{i},2)" if banco is not None else None))
    resumo["A9"] = (f"Valores em azul: resumo impresso pelo {ex.banco} {ex.origem_resumo}. "
                    "Diferença zero = planilha confere com o banco.")
    resumo["A9"].font = Font(size=9)
    resumo["A11"] = "Movimentação por tipo"
    resumo["A11"].font = Font(bold=True, size=11)
    cabecalho(resumo, 12, ["Tipo", "Entradas (R$)", "Saídas (R$)", "Líquido (R$)", "Qtde"])
    presentes = [t for t in TIPOS if any(m.tipo == t for m in ex.movimentos)]
    presentes += sorted({m.tipo for m in ex.movimentos} - set(presentes))
    tipo_col = f"Extrato!${L['tipo']}$6:${L['tipo']}${ultima}"
    r = 12
    for t in presentes:
        r += 1
        resumo.cell(r, 1, t).font = normal
        fmt(resumo.cell(r, 2, f"=SUMIFS(Extrato!${L['ent']}$6:${L['ent']}${ultima},{tipo_col},$A{r})"))
        fmt(resumo.cell(r, 3, f"=SUMIFS(Extrato!${L['sai']}$6:${L['sai']}${ultima},{tipo_col},$A{r})"))
        fmt(resumo.cell(r, 4, f"=B{r}-C{r}"))
        resumo.cell(r, 5, f"=COUNTIF({tipo_col},$A{r})").font = normal
    r += 1
    resumo.cell(r, 1, "TOTAL")
    for col, letra in ((2, "B"), (3, "C"), (4, "D"), (5, "E")):
        c = resumo.cell(r, col, f"=SUM({letra}13:{letra}{r - 1})")
        if col < 5:
            c.number_format = MOEDA
    for col in range(1, 6):
        resumo.cell(r, col).font = negrito
        resumo.cell(r, col).fill = cinza
    for col, w in zip("ABCDE", [34, 16, 16, 16, 8]):
        resumo.column_dimensions[col].width = w
    resumo.row_dimensions[3].height = resumo.row_dimensions[12].height = 20

    # ---- Investimentos CDB (Itaú)
    if ex.cdbs:
        inv = wb.create_sheet("Investimentos CDB")
        inv["A1"] = f"Carteira CDB e Renda Fixa — posição em {fim:%d/%m/%Y}" if fim else "Carteira CDB e Renda Fixa"
        inv["A1"].font = Font(bold=True, size=12)
        dia = f"{fim:%d/%m/%y}" if fim else ""
        cabecalho(inv, 3, ["Produto", "Data aplicação", "Vencimento", "Valor aplicado (R$)", "Remuneração (% DI)",
                           f"Saldo bruto {dia} (R$)", f"Saldo líquido {dia} (R$)", "Rendimento bruto (R$)",
                           "IR/IOF estimado (R$)"])
        r = 3
        for c in ex.cdbs:
            r += 1
            inv.cell(r, 1, c["produto"]).font = normal
            for col, chave in ((2, "aplicacao"), (3, "vencimento")):
                inv.cell(r, col, c[chave]).number_format = "dd/mm/yyyy"
                inv.cell(r, col).font = normal
            fmt(inv.cell(r, 4, c["valor"]))
            fmt(inv.cell(r, 5, c["taxa"]), "0.00%")
            fmt(inv.cell(r, 6, c["bruto"]))
            fmt(inv.cell(r, 7, c["liquido"]))
            fmt(inv.cell(r, 8, f"=F{r}-D{r}"))
            fmt(inv.cell(r, 9, f"=F{r}-G{r}"))
        ult = r
        r += 1
        inv.cell(r, 1, "TOTAL")
        for col, letra in ((4, "D"), (6, "F"), (7, "G"), (8, "H"), (9, "I")):
            fmt(inv.cell(r, col, f"=SUM({letra}4:{letra}{ult})"), fonte=negrito)
        for col in range(1, 10):
            inv.cell(r, col).fill = cinza
            if col < 4:
                inv.cell(r, col).font = negrito
        if ex.total_cdb:
            inv.cell(r + 2, 1, f"Totais do banco: aplicado {_br(ex.total_cdb[0])} · saldo bruto {_br(ex.total_cdb[1])} · "
                               f"saldo líquido {_br(ex.total_cdb[2])}.").font = Font(size=9)
        for col, w in zip("ABCDEFGHI", [11, 13, 12, 17, 13, 19, 19, 17, 17]):
            inv.column_dimensions[col].width = w
        inv.row_dimensions[3].height = 30
        inv.freeze_panes = "A4"
        inv.auto_filter.ref = f"A3:I{ult}"

    # ---- Saldos diários de investimento automático (Bradesco Invest Fácil)
    if ex.saldos_invest:
        inv = wb.create_sheet("Invest Fácil")
        inv["A1"] = "Saldos Invest Fácil / Plus"
        inv["A1"].font = Font(bold=True, size=12)
        inv["A2"] = "Saldo diário da aplicação automática informado pelo banco (separado do saldo da conta corrente)."
        inv["A2"].font = Font(size=9)
        cabecalho(inv, 4, ["Data", "Histórico", "Valor (R$)"])
        r = 4
        for d, hist, v in ex.saldos_invest:
            r += 1
            inv.cell(r, 1, d).number_format = "dd/mm/yyyy"
            inv.cell(r, 1).font = normal
            inv.cell(r, 2, hist).font = normal
            fmt(inv.cell(r, 3, v))
        for col, w in zip("ABC", [11, 26, 16]):
            inv.column_dimensions[col].width = w
        inv.freeze_panes = "A5"

    # ---- Lançamentos posteriores ao período
    if ex.ultimos:
        ws = wb.create_sheet("Últimos lançamentos")
        a, b = ex.ultimos[0].data, ex.ultimos[-1].data
        aba_extrato(ws, f"Últimos lançamentos (fora do período do extrato), de {a:%d/%m/%Y} a {b:%d/%m/%Y}",
                    ex.ultimos, ex.ultimos_saldo_anterior, ex.ultimos_data_anterior)

    wb.calculation = CalcProperties(fullCalcOnLoad=True)
    saida = BytesIO()
    wb.save(saida)
    return saida.getvalue()
