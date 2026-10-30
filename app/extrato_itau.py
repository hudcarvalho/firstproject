"""Extrato mensal do Itaú (conta corrente + aplicação automática) em planilha organizada.

Gera as abas:
- Resumo: saldos, entradas e saídas calculados x impressos pelo banco (diferença) e a
  movimentação por tipo (SOMASES / CONT.SE);
- Extrato: data, histórico, favorecido/remetente, tipo, entrada, saída, saldo por fórmula e,
  no fim de cada dia, o saldo do banco (conta corrente + aplicação automática) com "Confere";
- Investimentos CDB: carteira de CDB e renda fixa com rendimento e IR/IOF estimado.

As linhas internas da aplicação automática (Res/Apl Aplic Aut, SALDO APLIC AUT) não entram como
movimento: servem só para conferir o saldo do fim do dia.
"""
import re
import unicodedata
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO

MESES = {"jan": 1, "fev": 2, "mar": 3, "abr": 4, "mai": 5, "jun": 6, "jul": 7, "ago": 8,
         "set": 9, "out": 10, "nov": 11, "dez": 12}
NOMES_MES = ["janeiro", "fevereiro", "março", "abril", "maio", "junho", "julho", "agosto",
             "setembro", "outubro", "novembro", "dezembro"]
_VALOR = re.compile(r"^-?\d{1,3}(?:\.\d{3})*,\d{2}-?$")
_DIA = re.compile(r"^(\d{2})/(\d{2})$")
_CDB = re.compile(r"^(?P<produto>.+?)\s+(?P<ap>\d{2}/\d{2}/\d{2})\s+(?P<venc>\d{2}/\d{2}/\d{2})\s+"
                  r"(?P<valor>[\d.]+,\d{2})\s+(?P<taxa>[\d.]+,\d{2})%\s*DI\s+(?P<bruto>[\d.]+,\d{2})\s+"
                  r"(?P<liquido>[\d.]+,\d{2})$")

# Tipos na ordem em que aparecem no resumo.
TIPOS = [
    "Aplicação em CDB", "Resgate de CDB", "Recebimento PIX/TED", "Recebimento Sispag",
    "Cobrança (títulos liquidados)", "Depósitos", "Tributos/guias", "Transferência CC Itaú",
    "PIX enviado", "TED/DOC enviado", "Pagamento Sispag", "Débito automático",
    "Boletos/contas pagas", "Tarifas bancárias", "Juros/IOF", "Rendimento aplicação automática",
    "Outras entradas", "Outras saídas",
]


def _sem_acento(t):
    return "".join(c for c in unicodedata.normalize("NFD", t) if unicodedata.category(c) != "Mn").upper()


def _valor(t):
    neg = t.startswith("-") or t.endswith("-")
    v = Decimal(t.strip("-").replace(".", "").replace(",", "."))
    return -v if neg else v


@dataclass
class Movimento:
    data: date
    historico: str
    entrada: Decimal | None = None
    saida: Decimal | None = None
    saldo_banco: Decimal | None = None  # preenchido na última linha do dia
    favorecido: str = ""
    tipo: str = ""


@dataclass
class Extrato:
    empresa: str = ""
    agencia: str = ""
    conta: str = ""
    inicio: date | None = None
    fim: date | None = None
    saldo_anterior: Decimal | None = None
    saldo_final: Decimal | None = None
    total_entradas: Decimal | None = None
    total_saidas: Decimal | None = None
    tem_aplicacao_automatica: bool = False
    movimentos: list = field(default_factory=list)
    cdbs: list = field(default_factory=list)  # dicts
    total_cdb: tuple | None = None  # (aplicado, bruto, líquido) impressos pelo banco


def e_extrato_itau(pdf):
    """O PDF (pdfplumber aberto) é um extrato mensal do Itaú?"""
    if not pdf.pages:
        return False
    texto = pdf.pages[0].extract_text() or ""
    return bool(re.search(r"extrato mensal ag \d+ cc [\d-]+", texto)) and "Movimenta" in texto


# ---------------------------------------------------------------- leitura

def _linhas(page, x_min=0):
    """Palavras agrupadas em linhas (tolerância vertical de 3 pt), da esquerda para a direita."""
    linhas = []
    for w in sorted(page.extract_words(x_tolerance=1.5, y_tolerance=2), key=lambda w: (w["top"], w["x0"])):
        if w["x0"] < x_min:
            continue
        if linhas and abs(linhas[-1][0] - w["top"]) <= 3:
            linhas[-1][1].append(w)
        else:
            linhas.append([w["top"], [w]])
    return [sorted(ws, key=lambda w: w["x0"]) for _, ws in linhas]


def _cabecalho_mov(page):
    """Posições das colunas da movimentação na página (ou None se a página não tem a tabela)."""
    for ws in _linhas(page):
        # "entradas R$" às vezes vem colado ("entradasR$")
        textos = [w["text"][:-2] if w["text"].endswith("R$") and len(w["text"]) > 2 else w["text"] for w in ws]
        if "data" in textos and "descrição" in textos and "saldo" in textos and "entradas" in textos:
            def fim(rotulo):
                i = textos.index(rotulo)
                return ws[i + 1]["x1"] if i + 1 < len(ws) and ws[i + 1]["text"] == "R$" else ws[i]["x1"]
            return {"top": ws[0]["top"], "data": ws[textos.index("data")]["x0"],
                    "desc": ws[textos.index("descrição")]["x0"],
                    "entradas": fim("entradas"), "saidas": fim("saídas"), "saldo": fim("saldo")}
    return None


def _favorecido(hist):
    for prefixo in ("Sispag ", "BOLETO PAGO ", "PIX TRANSF ", "PIX ENVIADO ", "PIX RECEBIDO ", "TED ",
                    "DOC ", "DA ", "TBI ", "INT TRANSF ", "TRANSF "):
        if hist.upper().startswith(prefixo.upper()):
            resto = hist[len(prefixo):].strip()
            return re.sub(r"\s*\d{2}/\d{2}$", "", resto).strip()
    return ""


def _tipo(hist, entrada, debitos_automaticos):
    h = _sem_acento(hist)
    credito = entrada is not None
    if h.startswith("APLICACAO CDB") or h.startswith("APLIC CDB"):
        return "Aplicação em CDB"
    if h.startswith("RESGATE CDB") or h.startswith("RESG CDB"):
        return "Resgate de CDB"
    if h.startswith("REND PAGO APLIC"):
        return "Rendimento aplicação automática"
    if h.startswith("MOV TIT COB") or "COBRANCA" in h:
        return "Cobrança (títulos liquidados)"
    if h.startswith("DA ") or h in debitos_automaticos or h.startswith("SISDEB"):
        return "Débito automático"
    if h.startswith("SISPAG TRIB") or h.startswith("DARF") or h.startswith("GPS") or "TRIBUTO" in h:
        return "Tributos/guias"
    if h.startswith("SISPAG TRANSF CC ITAU") or h.startswith("TRANSF CC ITAU"):
        return "Transferência CC Itaú"
    if h.startswith("SISPAG CONCESSIONARIA") or h.startswith("BOLETO PAGO") or h.startswith("PAG BOLETO") \
            or h.startswith("PAGTO") or h.startswith("CONCESSIONARIA"):
        return "Boletos/contas pagas"
    if h.startswith("PIX ENVIADO"):
        return "PIX enviado"
    if h.startswith("PIX") or h.startswith("TED") or h.startswith("DOC") or h.startswith("TBI"):
        return "Recebimento PIX/TED" if credito else "TED/DOC enviado"
    if h.startswith("SISPAG"):
        return "Recebimento Sispag" if credito else "Pagamento Sispag"
    if h.startswith("DEP") or "DEPOSITO" in h:
        return "Depósitos"
    if h.startswith("TAR") or "TARIFA" in h:
        return "Tarifas bancárias"
    if h.startswith("IOF") or h.startswith("JUROS"):
        return "Juros/IOF"
    return "Outras entradas" if credito else "Outras saídas"


def ler(pdf):
    """Lê o extrato (pdfplumber aberto) e devolve um Extrato."""
    ex = Extrato()
    pagina1 = pdf.pages[0].extract_text() or ""
    m = re.search(r"extrato mensal ag (\d+) cc ([\d-]+) (\w{3}) (\d{4})", pagina1)
    if m:
        ex.agencia, ex.conta = m[1], m[2]
        mes, ano = MESES.get(m[3].lower(), 1), int(m[4])
        ex.inicio = date(ano, mes, 1)
    linhas1 = pagina1.splitlines()
    endereco = re.compile(r"^(AV|AVENIDA|R|RUA|AL|ALAMEDA|ROD|RODOVIA|ESTR|PCA|PRACA|TV|TRAVESSA|EST)\.?\s")
    for i, l in enumerate(linhas1[:15]):
        if i and endereco.match(l.strip().upper()):
            ex.empresa = linhas1[i - 1].strip()
            break
    m = re.search(r"saldo em (\d{2}/\d{2}/\d{2}) saldo em (\d{2}/\d{2}/\d{2})", pagina1)
    if m:
        ex.fim = datetime.strptime(m[2], "%d/%m/%y").date()
    m = re.search(r"R\$ (-?[\d.]+,\d{2}-?) R\$ (-?[\d.]+,\d{2}-?)\s*\n?\s*entradas", pagina1) or \
        re.search(r"Automáticas\s*\n\s*R\$ (-?[\d.]+,\d{2}-?) R\$ (-?[\d.]+,\d{2}-?)", pagina1)
    if m:
        ex.saldo_anterior, ex.saldo_final = _valor(m[1]), _valor(m[2])
    totais = re.findall(r"^total (-?[\d.]+,\d{2})$", pagina1, flags=re.M)
    if len(totais) >= 2:
        ex.total_entradas, ex.total_saidas = _valor(totais[0]), _valor(totais[1])

    # débitos automáticos efetuados (para classificar, ex.: "Business 4004-1849")
    debitos_automaticos = set()
    for page in pdf.pages:
        texto = page.extract_text() or ""
        if "Débitos automáticos efetuados" in texto:
            trecho = texto.split("Débitos automáticos efetuados", 1)[1].split("Débitos automáticos não efetuados")[0]
            for _, hist in re.findall(r"(\d{2}/\d{2}/\d{2}) (.+?) [\d.]+,\d{2}", trecho):
                debitos_automaticos.add(_sem_acento(hist.strip()))

    # movimentação
    dia = None
    saldo_aplic = None
    saldo_cc = Decimal("0")
    viu_saldo_no_dia = False
    fim = False
    movimentos_do_dia = []

    def fechar_dia():
        if dia is not None and viu_saldo_no_dia and movimentos_do_dia:
            movimentos_do_dia[-1].saldo_banco = (saldo_aplic or Decimal("0")) + saldo_cc \
                if ex.tem_aplicacao_automatica else saldo_cc

    for page in pdf.pages:
        if fim:
            break
        cab = _cabecalho_mov(page)
        if not cab:
            continue
        for ws in _linhas(page, x_min=cab["data"] - 3):
            if ws[0]["top"] <= cab["top"] + 2:
                continue
            textos = [w["text"] for w in ws]
            if textos[0] == "(créditos)":
                continue
            data_txt = None
            if _DIA.match(textos[0]) and ws[0]["x0"] < cab["desc"] - 5:
                data_txt = textos[0]
                ws, textos = ws[1:], textos[1:]
            valores, desc = [], []
            for w in ws:
                if _VALOR.match(w["text"]) and w["x0"] > cab["desc"] + 100:
                    col = min(("entradas", "saidas", "saldo"), key=lambda c: abs(cab[c] - w["x1"]))
                    valores.append((col, _valor(w["text"])))
                else:
                    desc.append(w["text"])
            historico = " ".join(desc).strip()
            if data_txt:
                novo = date(ex.inicio.year if ex.inicio else date.today().year,
                            int(data_txt[3:5]), int(data_txt[:2]))
                if dia is not None and novo != dia:
                    fechar_dia()
                    movimentos_do_dia = []
                    viu_saldo_no_dia = False
                dia = novo
            if not historico and not valores:
                continue
            h = _sem_acento(historico)
            saldo_col = next((v for c, v in valores if c == "saldo"), None)
            if h.startswith("SALDO ANTERIOR"):
                if ex.saldo_anterior is None and saldo_col is not None:
                    ex.saldo_anterior = saldo_col
                if ex.inicio is None and dia:
                    ex.inicio = dia
                saldo_anterior_data = dia
                continue
            if h.startswith("SALDO FINAL"):
                if saldo_col is not None and ex.saldo_final is None:
                    ex.saldo_final = saldo_col
                fim = True
                break
            if h.startswith("SALDO APLIC AUT") or h.startswith("SDO APLIC AUT"):
                ex.tem_aplicacao_automatica = True
                v = saldo_col if saldo_col is not None else next((v for _, v in valores), None)
                if v is not None:
                    saldo_aplic = v
                    viu_saldo_no_dia = True
                continue
            if h.startswith("SALDO EM C/C") or h.startswith("SALDO DO DIA") or h.startswith("SDO CTA"):
                if saldo_col is not None:
                    saldo_cc = saldo_col
                    viu_saldo_no_dia = True
                continue
            if saldo_col is not None:
                saldo_cc = saldo_col
                viu_saldo_no_dia = True
            if h.startswith("RES APLIC AUT") or h.startswith("APL APLIC AUT"):
                ex.tem_aplicacao_automatica = True
                continue
            entrada = next((v for c, v in valores if c == "entradas"), None)
            saida = next((v for c, v in valores if c == "saidas"), None)
            if entrada is None and saida is None:
                continue
            if saida is not None:
                saida = abs(saida)
            mov = Movimento(dia, historico, entrada, saida)
            mov.favorecido = _favorecido(historico)
            mov.tipo = _tipo(historico, entrada, debitos_automaticos)
            ex.movimentos.append(mov)
            movimentos_do_dia.append(mov)
    fechar_dia()
    if ex.inicio and ex.saldo_anterior is not None:
        ex.inicio = ex.inicio.replace(day=1)

    # carteira de CDB / renda fixa
    for page in pdf.pages:
        texto = page.extract_text() or ""
        if "CDB e Renda Fixa" not in texto:
            continue
        for linha in texto.splitlines():
            m = _CDB.match(linha.strip())
            if m:
                ex.cdbs.append({
                    "produto": m["produto"], "aplicacao": datetime.strptime(m["ap"], "%d/%m/%y").date(),
                    "vencimento": datetime.strptime(m["venc"], "%d/%m/%y").date(),
                    "valor": _valor(m["valor"]), "taxa": _valor(m["taxa"]) / 100,
                    "bruto": _valor(m["bruto"]), "liquido": _valor(m["liquido"])})
        m = re.search(r"total aplicações financeiras ([\d.]+,\d{2}) ([\d.]+,\d{2}) ([\d.]+,\d{2})", texto)
        if m:
            ex.total_cdb = tuple(_valor(x) for x in m.groups())
    return ex


# ---------------------------------------------------------------- planilha

def planilha(ex):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.workbook.properties import CalcProperties

    MOEDA = '#,##0.00;[Red]\\-#,##0.00;\\-'
    cab_fonte = Font(bold=True, color="FFFFFF", size=11)
    cab_fundo = PatternFill("solid", fgColor="1F3864")
    azul = Font(color="0000FF", size=10)
    normal = Font(size=10)
    cinza = PatternFill("solid", fgColor="F2F2F2")
    fim = ex.fim or (ex.movimentos[-1].data if ex.movimentos else None)
    inicio = ex.inicio
    mes_ano = f"{NOMES_MES[fim.month - 1]}/{fim.year}" if fim else ""
    conta = f"Itaú ag. {ex.agencia} c/c {ex.conta}"
    titulo = " — ".join(x for x in (ex.empresa, conta) if x)

    def cabecalho(ws, linha, textos):
        for i, t in enumerate(textos, start=1):
            c = ws.cell(linha, i, t)
            c.font, c.fill = cab_fonte, cab_fundo
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    def fmt(c, formato=MOEDA, fonte=normal):
        c.number_format = formato
        c.font = fonte
        return c

    wb = Workbook()
    resumo = wb.active
    resumo.title = "Resumo"
    ext = wb.create_sheet("Extrato")

    # ---- Extrato
    ext["A1"] = titulo
    ext["A1"].font = Font(bold=True, size=12)
    periodo = f"Extrato de {inicio:%d/%m/%Y} a {fim:%d/%m/%Y}" if inicio and fim else "Extrato"
    if ex.tem_aplicacao_automatica:
        periodo += " (saldo = conta corrente + aplicação automática Aplic Aut Mais)"
    ext["A2"] = periodo
    ext["A2"].font = normal
    cabecalho(ext, 4, ["Data", "Histórico (banco)", "Favorecido / Remetente", "Tipo", "Entrada (R$)",
                       "Saída (R$)", "Saldo (R$)", "Saldo banco fim do dia (R$)", "Confere"])
    data_anterior = None
    if inicio:
        from datetime import timedelta
        data_anterior = inicio - timedelta(days=1)
    ext["A5"] = data_anterior
    ext["B5"] = "SALDO ANTERIOR"
    ext["G5"] = ex.saldo_anterior
    for col in "ABCDEFGHI":
        c = ext[f"{col}5"]
        c.font, c.fill = Font(bold=True, size=10), cinza
    ext["A5"].number_format = "dd/mm/yyyy"
    ext["G5"].number_format = MOEDA
    ext["G5"].font = Font(bold=True, size=10, color="0000FF")
    r = 5
    for m in ex.movimentos:
        r += 1
        ext.cell(r, 1, m.data).number_format = "dd/mm/yyyy"
        ext.cell(r, 1).font = normal
        ext.cell(r, 2, m.historico).font = normal
        ext.cell(r, 3, m.favorecido or None).font = normal
        ext.cell(r, 4, m.tipo).font = normal
        fmt(ext.cell(r, 5, m.entrada))
        fmt(ext.cell(r, 6, m.saida))
        fmt(ext.cell(r, 7, f"=G{r - 1}+E{r}-F{r}"))
        if m.saldo_banco is not None:
            fmt(ext.cell(r, 8, m.saldo_banco), fonte=azul)
            ext.cell(r, 9, f'=IF(ROUND(G{r}-H{r},2)=0,"OK","DIVERGE")').font = normal
    ultima = r
    total = ultima + 1
    ext.cell(total, 2, "TOTAIS DO MÊS")
    fmt(ext.cell(total, 5, f"=SUM(E6:E{ultima})"))
    fmt(ext.cell(total, 6, f"=SUM(F6:F{ultima})"))
    fmt(ext.cell(total, 7, f"=G{ultima}"))
    for col in range(1, 10):
        ext.cell(total, col).font = Font(bold=True, size=10)
        ext.cell(total, col).fill = cinza
    for col, w in zip("ABCDEFGHI", [11, 32, 26, 30, 15, 15, 15, 16, 9]):
        ext.column_dimensions[col].width = w
    ext.row_dimensions[4].height = 30
    ext.freeze_panes = "A5"
    ext.auto_filter.ref = f"A4:I{ultima}"

    # ---- Resumo
    resumo["A1"] = f"{titulo} — {mes_ano}" if mes_ano else titulo
    resumo["A1"].font = Font(bold=True, size=12)
    cabecalho(resumo, 3, ["Item", "Calculado (R$)", "Banco (R$)", "Diferença"])
    itens = [
        (f"Saldo em {data_anterior:%d/%m/%y}" if data_anterior else "Saldo anterior", "=Extrato!G5", ex.saldo_anterior),
        ("Entradas (créditos)", f"=Extrato!E{total}", ex.total_entradas),
        ("Saídas (débitos)", f"=Extrato!F{total}", ex.total_saidas),
        (f"Saldo em {fim:%d/%m/%y}" if fim else "Saldo final", f"=Extrato!G{ultima}", ex.saldo_final),
    ]
    for i, (rotulo, formula, banco) in enumerate(itens, start=4):
        resumo.cell(i, 1, rotulo).font = normal
        fmt(resumo.cell(i, 2, formula))
        fmt(resumo.cell(i, 3, banco), fonte=azul)
        fmt(resumo.cell(i, 4, f"=ROUND(B{i}-C{i},2)" if banco is not None else None))
    resumo["A9"] = ("Valores em azul: resumo impresso pelo Itaú na 1ª página do extrato. "
                    "Diferença zero = planilha confere com o banco.")
    resumo["A9"].font = Font(size=9)
    resumo["A11"] = "Movimentação por tipo"
    resumo["A11"].font = Font(bold=True, size=11)
    cabecalho(resumo, 12, ["Tipo", "Entradas (R$)", "Saídas (R$)", "Líquido (R$)", "Qtde"])
    presentes = [t for t in TIPOS if any(m.tipo == t for m in ex.movimentos)]
    faixa = f"$6:$D${ultima}"
    r = 12
    for t in presentes:
        r += 1
        resumo.cell(r, 1, t).font = normal
        fmt(resumo.cell(r, 2, f"=SUMIFS(Extrato!$E$6:$E${ultima},Extrato!$D$6:$D${ultima},$A{r})"))
        fmt(resumo.cell(r, 3, f"=SUMIFS(Extrato!$F$6:$F${ultima},Extrato!$D$6:$D${ultima},$A{r})"))
        fmt(resumo.cell(r, 4, f"=B{r}-C{r}"))
        resumo.cell(r, 5, f"=COUNTIF(Extrato!$D$6:$D${ultima},$A{r})").font = normal
    r += 1
    resumo.cell(r, 1, "TOTAL")
    for col, letra in ((2, "B"), (3, "C"), (4, "D"), (5, "E")):
        c = resumo.cell(r, col, f"=SUM({letra}13:{letra}{r - 1})")
        if col < 5:
            c.number_format = MOEDA
    for col in range(1, 6):
        resumo.cell(r, col).font = Font(bold=True, size=10)
        resumo.cell(r, col).fill = cinza
    for col, w in zip("ABCDE", [34, 16, 16, 16, 8]):
        resumo.column_dimensions[col].width = w
    resumo.row_dimensions[3].height = resumo.row_dimensions[12].height = 20

    # ---- Investimentos CDB
    if ex.cdbs:
        inv = wb.create_sheet("Investimentos CDB")
        inv["A1"] = f"Carteira CDB e Renda Fixa — posição em {fim:%d/%m/%Y}" if fim else "Carteira CDB e Renda Fixa"
        inv["A1"].font = Font(bold=True, size=12)
        dia = f"{fim:%d/%m/%y}" if fim else ""
        cabecalho(inv, 3, ["Produto", "Data aplicação", "Vencimento", "Valor aplicado (R$)", "Remuneração (% DI)",
                           f"Saldo bruto {dia} (R$)", f"Saldo líquido {dia} (R$)", "Rendimento bruto (R$)",
                           "IR/IOF estimado (R$)"])
        r = 3
        for c in ex.cdbs:
            r += 1
            inv.cell(r, 1, c["produto"]).font = normal
            for col, chave in ((2, "aplicacao"), (3, "vencimento")):
                inv.cell(r, col, c[chave]).number_format = "dd/mm/yyyy"
                inv.cell(r, col).font = normal
            fmt(inv.cell(r, 4, c["valor"]))
            fmt(inv.cell(r, 5, c["taxa"]), "0.00%")
            fmt(inv.cell(r, 6, c["bruto"]))
            fmt(inv.cell(r, 7, c["liquido"]))
            fmt(inv.cell(r, 8, f"=F{r}-D{r}"))
            fmt(inv.cell(r, 9, f"=F{r}-G{r}"))
        ult = r
        r += 1
        inv.cell(r, 1, "TOTAL")
        for col, letra in ((4, "D"), (6, "F"), (7, "G"), (8, "H"), (9, "I")):
            fmt(inv.cell(r, col, f"=SUM({letra}4:{letra}{ult})"), fonte=Font(bold=True, size=10))
        for col in range(1, 10):
            inv.cell(r, col).fill = cinza
            if col < 4:
                inv.cell(r, col).font = Font(bold=True, size=10)
        if ex.total_cdb:
            def br(v):
                return f"{v:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
            inv.cell(r + 2, 1, f"Totais do banco: aplicado {br(ex.total_cdb[0])} · saldo bruto {br(ex.total_cdb[1])} · "
                               f"saldo líquido {br(ex.total_cdb[2])}.").font = Font(size=9)
        for col, w in zip("ABCDEFGHI", [11, 13, 12, 17, 13, 19, 19, 17, 17]):
            inv.column_dimensions[col].width = w
        inv.row_dimensions[3].height = 30
        inv.freeze_panes = "A4"
        inv.auto_filter.ref = f"A3:I{ult}"

    wb.calculation = CalcProperties(fullCalcOnLoad=True)
    saida = BytesIO()
    wb.save(saida)
    return saida.getvalue()
