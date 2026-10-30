"""Conversão de PDF em planilha Excel.

Lê o texto do PDF com o pdfplumber (só PDFs com texto; PDF escaneado, que é imagem, precisaria
de OCR) e monta as colunas de duas formas:
- "layout": usa o alinhamento do texto na página inteira para formar linhas e colunas (bom para
  relatórios contábeis sem bordas: balancetes, razões, DRE);
- "tabelas": só as tabelas desenhadas com linhas/bordas.
Valores no formato brasileiro (1.234,56, -1.234,56, (1.234,56), 1.234,56-) viram números.
"""
import re
from decimal import Decimal, InvalidOperation
from io import BytesIO

MAX_PAGINAS = 150
MODOS = {"auto": "Automático (reconhece o extrato mensal do Itaú; nos demais, layout da página)",
         "layout": "Layout da página (relatórios, com ou sem bordas)",
         "tabelas": "Somente tabelas com linhas/bordas"}

_NUMERO = re.compile(r"^(?P<neg1>[-(])?\s*(?:R\$\s*)?(?P<num>\d{1,3}(?:\.\d{3})*,\d{1,4}|\d+,\d{1,4})\s*(?P<neg2>[-)])?\s*(?P<dc>[DC])?$")


class ErroPDF(Exception):
    """PDF que não pôde ser lido (inválido, protegido por senha, sem texto…)."""


def numero_br(texto):
    """'1.234,56' -> Decimal; '(1.234,56)', '-1.234,56' e '1.234,56-' -> negativo. Senão None."""
    if not texto:
        return None
    m = _NUMERO.match(texto.strip())
    if not m:
        return None
    try:
        valor = Decimal(m["num"].replace(".", "").replace(",", "."))
    except InvalidOperation:
        return None
    return -valor if (m["neg1"] or m["neg2"]) else valor


def _limpar(celula):
    if celula is None:
        return ""
    return re.sub(r"\s+", " ", str(celula)).strip()


def _segmentos_por_linha(page):
    """Palavras da página agrupadas em linhas; em cada linha, palavras próximas formam um trecho.

    Devolve [[(x0, x1, texto), ...], ...] de cima para baixo.
    """
    palavras = page.extract_words(x_tolerance=1.5, y_tolerance=2.5, keep_blank_chars=False,
                                  use_text_flow=False, extra_attrs=["size"])
    linhas = []
    for w in sorted(palavras, key=lambda w: (round(w["top"]), w["x0"])):
        meio = (w["top"] + w["bottom"]) / 2
        if linhas and abs(linhas[-1]["meio"] - meio) <= max(2.0, w["size"] * 0.35):
            linhas[-1]["palavras"].append(w)
        else:
            linhas.append({"meio": meio, "palavras": [w]})
    resultado = []
    for linha in linhas:
        segmentos = []
        for w in sorted(linha["palavras"], key=lambda w: w["x0"]):
            distancia = max(3.5, w["size"] * 0.75)  # espaço maior que ~3/4 de letra separa colunas
            if segmentos and w["x0"] - segmentos[-1][1] <= distancia:
                x0, _, texto = segmentos[-1]
                segmentos[-1] = (x0, w["x1"], texto + " " + w["text"])
            else:
                segmentos.append((w["x0"], w["x1"], w["text"]))
        resultado.append(segmentos)
    return resultado


def _faixas_de_colunas(linhas, largura_pagina):
    """Colunas = faixas horizontais ocupadas por texto nas linhas com mais de um trecho; os
    vãos verticais vazios ao longo da página separam uma coluna da outra."""
    multiplas = [l for l in linhas if len(l) > 1]
    if not multiplas:
        return []
    # as colunas vêm das linhas com cara de linha de tabela (quantidade típica de trechos);
    # títulos e rodapés com só dois pedaços não apagam os vãos entre colunas
    quantidades = sorted(len(l) for l in multiplas)
    mediana = quantidades[len(quantidades) // 2]
    base = [l for l in multiplas if len(l) >= mediana]
    if len(base) >= 2:
        multiplas = base
    passos = int(largura_pagina) + 2
    cobertura = [0] * passos
    for l in multiplas:
        for x0, x1, _ in l:
            if x1 - x0 > largura_pagina * 0.45:  # títulos e textos corridos não definem colunas
                continue
            for x in range(max(0, int(x0)), min(passos, int(x1) + 1)):
                cobertura[x] += 1
    # ignora poucas linhas fora do padrão (títulos, totais longos): ~5%, e ao menos 1 em tabelas maiores
    tolerancia = max(1 if len(multiplas) >= 8 else 0, len(multiplas) // 20)
    faixas, inicio = [], None
    for x, c in enumerate(cobertura + [0]):
        if c > tolerancia and inicio is None:
            inicio = x
        elif c <= tolerancia and inicio is not None:
            faixas.append((inicio, x))
            inicio = None
    return faixas


def _linhas_por_texto(page):
    """Todo o texto da página em linhas e colunas pelo alinhamento das palavras."""
    linhas = _segmentos_por_linha(page)
    faixas = _faixas_de_colunas(linhas, page.width)
    if len(faixas) < 2:
        return [[" ".join(t for _, _, t in l)] for l in linhas]

    def coluna(x):
        for i, (a, b) in enumerate(faixas):
            if a <= x <= b:
                return i
        return min(range(len(faixas)), key=lambda i: min(abs(x - faixas[i][0]), abs(x - faixas[i][1])))

    tabela = []
    for l in linhas:
        celulas = [""] * len(faixas)
        for x0, x1, texto in l:
            # linha com um só trecho (título, observação): vale onde o texto começa;
            # nas demais, o centro do trecho (valores costumam ser alinhados à direita)
            i = coluna(x0 if len(l) == 1 else (x0 + x1) / 2)
            celulas[i] = (celulas[i] + " " + texto).strip()
        tabela.append(celulas)
    return tabela


def _linhas_por_tabelas(page):
    linhas = []
    for tabela in page.extract_tables():
        if linhas:
            linhas.append([])  # linha em branco entre tabelas da mesma página
        linhas.extend(tabela)
    return linhas


def _sem_colunas_vazias(linhas):
    """Remove linhas vazias e colunas que ficaram vazias em todas as linhas."""
    linhas = [[_limpar(c) for c in l] for l in linhas]
    linhas = [l for l in linhas if any(l)]
    if not linhas:
        return []
    largura = max(len(l) for l in linhas)
    linhas = [l + [""] * (largura - len(l)) for l in linhas]
    usadas = [i for i in range(largura) if any(l[i] for l in linhas)]
    return [[l[i] for i in usadas] for l in linhas]


def converter(conteudo, modo="auto", aba_por_pagina=False, converter_numeros=True):
    """PDF (bytes) -> (planilha .xlsx em bytes, nº de páginas, nº de linhas).

    No modo "auto", extrato mensal do Itaú vira a planilha organizada de app.extrato_itau
    (resumo conferido com o banco, extrato classificado e investimentos)."""
    import pdfplumber
    from openpyxl import Workbook
    from openpyxl.styles import Font

    try:
        pdf = pdfplumber.open(BytesIO(conteudo))
        paginas = pdf.pages
        total = len(paginas)
    except Exception as e:  # PDF inválido, corrompido ou protegido por senha
        raise ErroPDF("Não foi possível ler o arquivo. Confira se é um PDF válido e sem senha.") from e
    if total > MAX_PAGINAS:
        pdf.close()
        raise ErroPDF(f"O PDF tem {total} páginas; o limite é {MAX_PAGINAS}. Divida o arquivo e converta por partes.")

    if modo == "auto":
        from . import extrato_itau
        try:
            reconhecido = extrato_itau.e_extrato_itau(pdf)
        except Exception:
            reconhecido = False
        if reconhecido:
            try:
                ex = extrato_itau.ler(pdf)
            finally:
                pdf.close()
            if ex.movimentos:
                return extrato_itau.planilha(ex), total, len(ex.movimentos)
            pdf = pdfplumber.open(BytesIO(conteudo))
            paginas = pdf.pages
    extrair = _linhas_por_tabelas if modo == "tabelas" else _linhas_por_texto
    wb = Workbook()
    ws = wb.active
    ws.title = "Página 1" if aba_por_pagina else "PDF"
    total_linhas = 0
    larguras = {}
    try:
        for n, page in enumerate(paginas, start=1):
            if aba_por_pagina and n > 1:
                ws = wb.create_sheet(f"Página {n}")
            linhas = _sem_colunas_vazias(extrair(page))
            if not aba_por_pagina and linhas and ws.max_row > 1:
                ws.append([])
            for linha in linhas:
                valores = []
                for c in linha:
                    v = numero_br(c) if converter_numeros else None
                    valores.append(v if v is not None else c)
                ws.append(valores)
                r = ws.max_row
                for i, v in enumerate(valores, start=1):
                    if isinstance(v, Decimal):
                        ws.cell(r, i).number_format = "#,##0.00;-#,##0.00"
                    larguras[(ws.title, i)] = max(larguras.get((ws.title, i), 0), len(str(v)))
                total_linhas += 1
            page.flush_cache()
    finally:
        pdf.close()
    if total_linhas == 0:
        raise ErroPDF("Nenhum texto encontrado no PDF. Se for um documento escaneado (imagem), "
                      "ele precisa passar por OCR antes; se escolheu “Somente tabelas”, tente “Layout da página”.")
    from openpyxl.utils import get_column_letter
    for (aba, col), w in larguras.items():
        wb[aba].column_dimensions[get_column_letter(col)].width = min(max(w + 2, 8), 60)
    for planilha in wb.worksheets:
        planilha.freeze_panes = None
        if planilha.max_row == 1 and planilha.cell(1, 1).value is None:
            planilha.cell(1, 1, "(página sem texto)").font = Font(italic=True, color="888888")
    saida = BytesIO()
    wb.save(saida)
    return saida.getvalue(), total, total_linhas
