"""Aba LALUR (empresas do Lucro Real).

Parte A: por período (trimestre ou anual), lucro líquido antes do IRPJ/CSLL, adições e
exclusões, compensação de prejuízo fiscal / base negativa (limite de 30% do lucro ajustado)
e cálculo do IRPJ (15% + adicional de 10% sobre o que passar de R$ 20.000,00 por mês) e da
CSLL. Parte B: saldos de prejuízo fiscal e de base negativa, encadeados de um período ao
outro a partir do saldo de abertura informado.
"""
from dataclasses import dataclass, field
from datetime import date
from decimal import ROUND_HALF_UP, Decimal

from flask import Blueprint, abort, flash, redirect, render_template, request, url_for

from .competencia import parse_valor
from .models import Empresa, LalurApuracao, LalurLancamento, LalurSaldoInicial, db

bp = Blueprint("lalur", __name__, url_prefix="/lalur")

ZERO = Decimal("0")
LIMITE_COMPENSACAO = Decimal("0.30")
ALIQUOTA_IRPJ = Decimal("0.15")
ALIQUOTA_ADICIONAL = Decimal("0.10")
LIMITE_ADICIONAL_MES = Decimal("20000")

SUGESTOES_ADICAO = [
    "Multas por infrações fiscais", "Brindes", "Doações não dedutíveis",
    "Provisões não dedutíveis", "Despesas com alimentação de sócios/dirigentes",
    "Resultado negativo de equivalência patrimonial", "Gratificações a administradores",
    "Despesas não necessárias à atividade", "Depreciação acima do limite fiscal",
]
SUGESTOES_EXCLUSAO = [
    "Resultado positivo de equivalência patrimonial", "Dividendos recebidos",
    "Reversão de provisões adicionadas anteriormente", "Juros sobre o capital próprio pagos",
    "Depreciação acelerada incentivada", "Ganhos de AVJ (diferidos)",
]


# Linhas do demonstrativo: (rótulo, valor no Demonstrativo, estilo). Estilo "secao" = título.
LINHAS = [
    ("Lucro líquido antes do IRPJ/CSLL", lambda d: d.apuracao.lucro_antes, ""),
    ("IRPJ", None, "secao"),
    ("(+) Adições", lambda d: d.irpj.adicoes, ""),
    ("(−) Exclusões", lambda d: d.irpj.exclusoes, ""),
    ("= Lucro ajustado", lambda d: d.irpj.ajustado, "subtotal"),
    ("(−) Compensação de prejuízo fiscal", lambda d: d.irpj.compensacao, ""),
    ("= Lucro real", lambda d: d.irpj.base, "subtotal"),
    ("IRPJ 15%", lambda d: d.irpj_normal, ""),
    ("Adicional 10%", lambda d: d.irpj_adicional, ""),
    ("(−) Deduções (retenções/antecipações)", lambda d: d.apuracao.irpj_deduzir, ""),
    ("IRPJ a pagar", lambda d: d.irpj_a_pagar, "total"),
    ("CSLL", None, "secao"),
    ("(+) Adições", lambda d: d.csll.adicoes, ""),
    ("(−) Exclusões", lambda d: d.csll.exclusoes, ""),
    ("= Base ajustada", lambda d: d.csll.ajustado, "subtotal"),
    ("(−) Compensação de base negativa", lambda d: d.csll.compensacao, ""),
    ("= Base de cálculo", lambda d: d.csll.base, "subtotal"),
    ("CSLL devida", lambda d: d.csll_devida, ""),
    ("(−) Deduções (retenções/antecipações)", lambda d: d.apuracao.csll_deduzir, ""),
    ("CSLL a pagar", lambda d: d.csll_a_pagar, "total"),
    ("Parte B — saldos ao final do período", None, "secao"),
    ("Prejuízo fiscal a compensar (IRPJ)", lambda d: d.irpj.saldo_final, ""),
    ("Base negativa a compensar (CSLL)", lambda d: d.csll.saldo_final, ""),
    ("Resultado do período", None, "secao"),
    ("Lucro líquido antes do IRPJ/CSLL", lambda d: d.apuracao.lucro_antes, ""),
    ("(−) IRPJ devido", lambda d: d.irpj_devido, ""),
    ("(−) CSLL devida", lambda d: d.csll_devida, ""),
    ("= Lucro líquido do período", lambda d: d.lucro_liquido, "total"),
]


def rotulo_periodo(periodo, curto=False):
    if periodo == 0:
        return "Anual"
    return f"{periodo}º trim." if curto else f"{periodo}º trimestre"


def _q(v):
    return v.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)


@dataclass
class Tributo:
    """Apuração de um tributo (IRPJ ou CSLL) em um período."""

    adicoes: Decimal = ZERO
    exclusoes: Decimal = ZERO
    ajustado: Decimal = ZERO        # lucro ajustado antes da compensação
    saldo_anterior: Decimal = ZERO  # prejuízo fiscal / base negativa disponível
    compensacao: Decimal = ZERO
    gerado: Decimal = ZERO          # prejuízo / base negativa gerada no período
    base: Decimal = ZERO            # lucro real / base de cálculo da CSLL

    @property
    def saldo_final(self):
        return self.saldo_anterior - self.compensacao + self.gerado


@dataclass
class Demonstrativo:
    apuracao: LalurApuracao
    irpj: Tributo = field(default_factory=Tributo)
    csll: Tributo = field(default_factory=Tributo)
    irpj_normal: Decimal = ZERO
    irpj_adicional: Decimal = ZERO
    csll_devida: Decimal = ZERO

    @property
    def irpj_devido(self):
        return self.irpj_normal + self.irpj_adicional

    @property
    def irpj_a_pagar(self):
        return max(self.irpj_devido - self.apuracao.irpj_deduzir, ZERO)

    @property
    def lucro_liquido(self):
        """Lucro líquido do período depois do IRPJ e da CSLL devidos."""
        return self.apuracao.lucro_antes - self.irpj_devido - self.csll_devida

    @property
    def csll_a_pagar(self):
        return max(self.csll_devida - self.apuracao.csll_deduzir, ZERO)


def _apurar_tributo(apuracao, chave, saldo):
    t = Tributo(saldo_anterior=saldo)
    for l in apuracao.lancamentos:
        if l.tributo in ("ambos", chave):
            if l.tipo == "adicao":
                t.adicoes += l.valor
            else:
                t.exclusoes += l.valor
    t.ajustado = apuracao.lucro_antes + t.adicoes - t.exclusoes
    if t.ajustado < 0:
        t.gerado = -t.ajustado
    else:
        if apuracao.compensar:
            t.compensacao = min(saldo, _q(t.ajustado * LIMITE_COMPENSACAO))
        t.base = t.ajustado - t.compensacao
    return t


def ordem(apuracao):
    """Ordem cronológica: trimestres 1-4 e a anual (0) ao fim do ano."""
    return apuracao.ano, apuracao.periodo or 5


def demonstrativos(empresa):
    """Calcula todas as apurações da empresa em ordem, encadeando os saldos da Parte B.

    Devolve {(ano, periodo): Demonstrativo} e os saldos finais (prejuízo, base negativa).
    Períodos sem lucro informado não entram no cálculo (os saldos passam direto).
    """
    abertura = db.session.query(LalurSaldoInicial).filter_by(empresa_id=empresa.id).first()
    prejuizo = abertura.prejuizo_fiscal if abertura else ZERO
    base_negativa = abertura.base_negativa if abertura else ZERO
    resultado = {}
    apuracoes = LalurApuracao.query.filter_by(empresa_id=empresa.id).all()
    for a in sorted(apuracoes, key=ordem):
        if a.lucro_antes is None:
            continue
        d = Demonstrativo(a)
        d.irpj = _apurar_tributo(a, "irpj", prejuizo)
        d.csll = _apurar_tributo(a, "csll", base_negativa)
        meses = 12 if a.periodo == 0 else 3
        d.irpj_normal = _q(d.irpj.base * ALIQUOTA_IRPJ)
        d.irpj_adicional = _q(max(d.irpj.base - LIMITE_ADICIONAL_MES * meses, ZERO) * ALIQUOTA_ADICIONAL)
        d.csll_devida = _q(d.csll.base * a.aliquota_csll / 100)
        prejuizo, base_negativa = d.irpj.saldo_final, d.csll.saldo_final
        resultado[(a.ano, a.periodo)] = d
    return resultado, (prejuizo, base_negativa)


# ---------------------------------------------------------------- telas

def _empresa(id_):
    empresa = db.session.get(Empresa, id_)
    if empresa is None:
        abort(404)
    return empresa


def _ano():
    try:
        ano = int(request.values.get("ano", ""))
    except ValueError:
        return date.today().year
    return ano if 2000 <= ano <= 2100 else date.today().year


def _exige_lucro_real(empresa):
    if not empresa.lucro_real:
        flash("O LALUR é usado só por empresas do Lucro Real.", "warning")
        return redirect(url_for("main.empresa_detalhe", id=empresa.id))
    return None


def _periodo(valor):
    if valor not in (0, 1, 2, 3, 4):
        abort(404)
    return valor


def _modo(empresa, ano):
    """'anual' ou 'trimestral' para o ano (pelo que já foi lançado; padrão trimestral)."""
    periodos = {p for (p,) in db.session.query(LalurApuracao.periodo)
                .filter_by(empresa_id=empresa.id, ano=ano)}
    if 0 in periodos:
        return "anual", True
    if periodos:
        return "trimestral", True
    return ("anual" if request.args.get("modo") == "anual" else "trimestral"), False


@bp.route("/<int:id>")
def empresa(id):
    empresa = _empresa(id)
    if (r := _exige_lucro_real(empresa)):
        return r
    ano = _ano()
    modo, modo_fixo = _modo(empresa, ano)
    periodos = [0] if modo == "anual" else [1, 2, 3, 4]
    calc, _ = demonstrativos(empresa)
    existentes = {a.periodo: a for a in LalurApuracao.query.filter_by(empresa_id=empresa.id, ano=ano)}
    colunas = [{"periodo": p, "rotulo": rotulo_periodo(p, curto=True),
                "apuracao": existentes.get(p), "d": calc.get((ano, p))} for p in periodos]
    abertura = LalurSaldoInicial.query.filter_by(empresa_id=empresa.id).first()
    return render_template("lalur/empresa.html", empresa=empresa, ano=ano, modo=modo,
                           modo_fixo=modo_fixo, colunas=colunas, abertura=abertura, linhas=LINHAS,
                           apurados=[c for c in colunas if c["d"]])


@bp.post("/<int:id>/saldo-inicial")
def saldo_inicial(id):
    empresa = _empresa(id)
    try:
        prejuizo = parse_valor(request.form.get("prejuizo_fiscal")) or ZERO
        base = parse_valor(request.form.get("base_negativa")) or ZERO
        if prejuizo < 0 or base < 0:
            raise ValueError
    except ValueError:
        flash("Saldos de abertura inválidos: informe valores positivos, como 15.000,00.", "danger")
    else:
        abertura = LalurSaldoInicial.query.filter_by(empresa_id=empresa.id).first()
        if abertura is None:
            abertura = LalurSaldoInicial(empresa_id=empresa.id)
            db.session.add(abertura)
        abertura.prejuizo_fiscal, abertura.base_negativa = prejuizo, base
        db.session.commit()
        flash("Saldos de abertura da Parte B salvos.", "success")
    return redirect(url_for("lalur.empresa", id=id, ano=_ano()))


def _obter_apuracao(empresa, ano, periodo, criar=False):
    a = LalurApuracao.query.filter_by(empresa_id=empresa.id, ano=ano, periodo=periodo).first()
    if a is None and criar:
        a = LalurApuracao(empresa_id=empresa.id, ano=ano, periodo=periodo, compensar=True,
                          aliquota_csll=Decimal("9"), irpj_deduzir=ZERO, csll_deduzir=ZERO)
        db.session.add(a)
    return a


def _conflito_de_modo(empresa, ano, periodo):
    """Não mistura apuração anual e trimestral no mesmo ano."""
    outros = LalurApuracao.query.filter_by(empresa_id=empresa.id, ano=ano)
    if periodo == 0:
        return outros.filter(LalurApuracao.periodo != 0).count() > 0
    return outros.filter_by(periodo=0).count() > 0


@bp.route("/<int:id>/<int:ano>/<int:periodo>")
def periodo(id, ano, periodo):
    empresa = _empresa(id)
    if (r := _exige_lucro_real(empresa)):
        return r
    periodo = _periodo(periodo)
    apuracao = _obter_apuracao(empresa, ano, periodo)
    calc, _ = demonstrativos(empresa)
    rotulo = rotulo_periodo(periodo)
    return render_template(
        "lalur/periodo.html", empresa=empresa, ano=ano, periodo=periodo, rotulo=rotulo,
        apuracao=apuracao, d=calc.get((ano, periodo)), tipos=LalurLancamento.TIPOS,
        tributos=LalurLancamento.TRIBUTOS, sugestoes_adicao=SUGESTOES_ADICAO,
        sugestoes_exclusao=SUGESTOES_EXCLUSAO,
    )


def _voltar_periodo(id, ano, periodo, ancora=""):
    return redirect(url_for("lalur.periodo", id=id, ano=ano, periodo=periodo) + ancora)


@bp.post("/<int:id>/<int:ano>/<int:periodo>")
def salvar(id, ano, periodo):
    empresa = _empresa(id)
    periodo = _periodo(periodo)
    if _conflito_de_modo(empresa, ano, periodo):
        flash("Este ano já tem apuração em outro formato (anual × trimestral).", "danger")
        return redirect(url_for("lalur.empresa", id=id, ano=ano))
    f = request.form
    try:
        lucro = parse_valor(f.get("lucro_antes"))
        aliquota = parse_valor(f.get("aliquota_csll")) or Decimal("9")
        irpj_ded = parse_valor(f.get("irpj_deduzir")) or ZERO
        csll_ded = parse_valor(f.get("csll_deduzir")) or ZERO
        if not ZERO <= aliquota <= 100 or irpj_ded < 0 or csll_ded < 0:
            raise ValueError
    except ValueError:
        flash("Valores inválidos. Use o formato 1.234,56 (o lucro pode ser negativo: -1.234,56).", "danger")
        return _voltar_periodo(id, ano, periodo)
    a = _obter_apuracao(empresa, ano, periodo, criar=True)
    a.lucro_antes = lucro
    a.compensar = f.get("compensar") == "on"
    a.aliquota_csll, a.irpj_deduzir, a.csll_deduzir = aliquota, irpj_ded, csll_ded
    a.observacao = (f.get("observacao") or "").strip() or None
    db.session.commit()
    flash("Apuração salva." if lucro is not None else
          "Apuração salva. Informe o lucro antes do IRPJ/CSLL para calcular.", "success")
    return _voltar_periodo(id, ano, periodo)


@bp.post("/<int:id>/<int:ano>/<int:periodo>/lancamentos")
def lancamento_novo(id, ano, periodo):
    empresa = _empresa(id)
    periodo = _periodo(periodo)
    if _conflito_de_modo(empresa, ano, periodo):
        flash("Este ano já tem apuração em outro formato (anual × trimestral).", "danger")
        return redirect(url_for("lalur.empresa", id=id, ano=ano))
    f = request.form
    tipo, tributo = f.get("tipo"), f.get("tributo") or "ambos"
    descricao = (f.get("descricao") or "").strip()
    try:
        valor = parse_valor(f.get("valor"))
    except ValueError:
        valor = None
    if tipo not in LalurLancamento.TIPOS or tributo not in LalurLancamento.TRIBUTOS \
            or not descricao or valor is None or valor <= 0:
        flash("Lançamento inválido: informe tipo, descrição e um valor positivo.", "danger")
        return _voltar_periodo(id, ano, periodo, "#lancamentos")
    a = _obter_apuracao(empresa, ano, periodo, criar=True)
    a.lancamentos.append(LalurLancamento(tipo=tipo, tributo=tributo, descricao=descricao[:200],
                                         valor=valor))
    db.session.commit()
    flash(f"{LalurLancamento.TIPOS[tipo]} lançada.", "success")
    return _voltar_periodo(id, ano, periodo, "#lancamentos")


@bp.post("/lancamentos/<int:lid>/remover")
def lancamento_remover(lid):
    l = db.session.get(LalurLancamento, lid)
    if l is None:
        abort(404)
    a = l.apuracao
    db.session.delete(l)
    db.session.commit()
    flash("Lançamento removido.", "success")
    return _voltar_periodo(a.empresa_id, a.ano, a.periodo, "#lancamentos")


@bp.post("/<int:id>/<int:ano>/<int:periodo>/excluir")
def excluir(id, ano, periodo):
    empresa = _empresa(id)
    a = _obter_apuracao(empresa, ano, _periodo(periodo))
    if a is not None:
        db.session.delete(a)
        db.session.commit()
        flash("Apuração do período excluída.", "success")
    return redirect(url_for("lalur.empresa", id=id, ano=ano))


# ---------------------------------------------------------------- exportação (PDF e Excel)

def _fmt(v):
    from .competencia import fmt_valor
    return fmt_valor(v) if v is not None else ""


def _nome_arquivo(empresa, ano, periodos, extensao):
    sufixo = "todos" if len(periodos) > 1 else ("anual" if periodos[0][0] == 0 else f"{periodos[0][0]}tri")
    return f"lalur_{empresa.codigo or empresa.id}_{ano}_{sufixo}.{extensao}"


def _pdf(empresa, ano, periodos):
    """periodos: [(periodo, Demonstrativo)] já apurados."""
    from .pdf_simples import Documento

    doc = Documento(f"LALUR {ano} - {empresa.razao_social}")
    margem, largura_util = 36, doc.largura_pagina - 72
    direita = margem + largura_util
    y = 0
    pagina = [0]

    def cabecalho():
        nonlocal y
        pagina[0] += 1
        doc.texto(margem, 40, "LALUR — Livro de Apuração do Lucro Real", 13, negrito=True)
        doc.texto(direita, 40, f"Exercício {ano}", 11, negrito=True, alinhar="direita")
        doc.texto(margem, 56, empresa.razao_social, 10)
        info = " · ".join(x for x in [f"Cód. {empresa.codigo}" if empresa.codigo else "",
                                       f"CNPJ {empresa.cnpj_formatado}", empresa.regime.nome] if x)
        doc.texto(margem, 69, info, 8.5, cinza=0.35)
        doc.linha(margem, 76, direita, 76, 1, 0.3)
        doc.texto(margem, doc.altura_pagina - 24,
                  f"Emitido em {date.today():%d/%m/%Y} pelo Controle de Rotina", 7.5, cinza=0.45)
        doc.texto(direita, doc.altura_pagina - 24, f"Página {pagina[0]}", 7.5, alinhar="direita", cinza=0.45)
        y = 96

    def espaco(altura):
        nonlocal y
        if y + altura > doc.altura_pagina - 44:
            doc.nova_pagina()
            cabecalho()

    def titulo(texto):
        nonlocal y
        espaco(40)
        doc.texto(margem, y, texto, 11, negrito=True)
        y += 10

    def tabela(cabecalhos, larguras, linhas, alinhar_numeros=True):
        """linhas: [(celulas, estilo)]; 1ª coluna texto, demais valores à direita."""
        nonlocal y
        alt = 14

        def topo():
            nonlocal y
            doc.faixa(margem, y, largura_util, alt, 0.85)
            x = margem
            for i, (c, w) in enumerate(zip(cabecalhos, larguras)):
                if i == 0 or not alinhar_numeros:
                    doc.texto(x + 4, y + 10, c, 8.5, negrito=True)
                else:
                    doc.texto(x + w - 4, y + 10, c, 8.5, negrito=True, alinhar="direita")
                x += w
            y += alt

        espaco(alt * 3)
        topo()
        for celulas, estilo in linhas:
            if y + alt > doc.altura_pagina - 44:
                doc.nova_pagina()
                cabecalho()
                topo()
            if estilo == "secao":
                doc.faixa(margem, y, largura_util, alt, 0.95)
                doc.texto(margem + 4, y + 10, celulas[0].upper(), 7.5, negrito=True, cinza=0.35)
            else:
                negrito = estilo in ("subtotal", "total")
                if estilo == "total":
                    doc.faixa(margem, y, largura_util, alt, 0.97)
                x = margem
                for i, (c, w) in enumerate(zip(celulas, larguras)):
                    if i == 0 or not alinhar_numeros:
                        doc.texto(x + 4, y + 10, c, 8.5, negrito=negrito)
                    else:
                        doc.texto(x + w - 4, y + 10, c, 8.5, negrito=negrito, alinhar="direita")
                    x += w
                doc.linha(margem, y + alt, direita, y + alt, 0.3, 0.8)
            y += alt
        y += 14

    cabecalho()
    # Parte A: uma coluna por período
    rot = largura_util * (0.42 if len(periodos) > 1 else 0.65)
    col = (largura_util - rot) / len(periodos)
    titulo("Parte A — Demonstração do lucro real e cálculo do IRPJ e da CSLL")
    linhas = []
    for rotulo, valor, estilo in LINHAS:
        if estilo == "secao":
            linhas.append(([rotulo], "secao"))
        else:
            linhas.append(([rotulo] + [_fmt(valor(d)) for _, d in periodos], estilo))
    tabela(["R$"] + [rotulo_periodo(p, curto=True) for p, _ in periodos],
           [rot] + [col] * len(periodos), linhas)

    # Adições e exclusões de cada período
    for p, d in periodos:
        itens = d.apuracao.lancamentos
        titulo(f"Adições e exclusões — {rotulo_periodo(p)}")
        if not itens:
            espaco(16)
            doc.texto(margem, y + 6, "Nenhuma adição ou exclusão lançada.", 8.5, cinza=0.4)
            y += 24
        else:
            larguras = [70, largura_util - 70 - 90 - 90, 90, 90]
            tabela(["Tipo", "Descrição", "Tributo", "Valor (R$)"], larguras,
                   [([LalurLancamento.TIPOS[l.tipo], l.descricao[:80],
                      LalurLancamento.TRIBUTOS[l.tributo], _fmt(l.valor)], "") for l in itens],
                   alinhar_numeros=False)
        if d.apuracao.observacao:
            espaco(16)
            doc.texto(margem, y + 6, f"Observação: {d.apuracao.observacao[:140]}", 8.5, cinza=0.3)
            y += 22
    return doc.gerar()


def _excel(empresa, ano, periodos):
    """Planilha com fórmulas: os dados digitados entram como valores e todo o cálculo do LALUR
    (adições, exclusões, compensação, IRPJ, adicional, CSLL, saldos da Parte B e lucro líquido)
    é feito por fórmulas do Excel, com a explicação de cada uma na última coluna."""
    from io import BytesIO

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.workbook.properties import CalcProperties

    negrito = Font(bold=True)
    fundo = PatternFill("solid", fgColor="D9E1F2")
    secao = PatternFill("solid", fgColor="F2F2F2")
    entrada = PatternFill("solid", fgColor="FFF2CC")  # células digitáveis
    numero = "#,##0.00;-#,##0.00"
    cinza = Font(color="666666", italic=True)

    def cabecalho(ws):
        ws.append(["LALUR — Livro de Apuração do Lucro Real", None, f"Exercício {ano}"])
        ws["A1"].font = Font(bold=True, size=13)
        ws.append([empresa.razao_social])
        ws.append([f"Cód. {empresa.codigo or '-'} · CNPJ {empresa.cnpj_formatado} · {empresa.regime.nome}"])
        ws.append([])

    wb = Workbook()
    ws = wb.active
    ws.title = "Parte A"

    # ---- Parâmetros (referenciados pelas fórmulas)
    par = wb.create_sheet("Parâmetros")
    parametros = [
        ("Alíquota do IRPJ", ALIQUOTA_IRPJ, "0.00%"),
        ("Alíquota do adicional de IRPJ", ALIQUOTA_ADICIONAL, "0.00%"),
        ("Parcela isenta do adicional por mês (R$)", LIMITE_ADICIONAL_MES, numero),
        ("Limite de compensação de prejuízo / base negativa", LIMITE_COMPENSACAO, "0.00%"),
    ]
    par.append(["Parâmetro", "Valor"])
    for c in par[1]:
        c.font, c.fill = negrito, fundo
    for nome, valor, formato in parametros:
        par.append([nome, valor])
        par.cell(par.max_row, 2).number_format = formato
        par.cell(par.max_row, 2).fill = entrada
    par.column_dimensions["A"].width = 52
    par.column_dimensions["B"].width = 16
    P_IRPJ, P_ADIC, P_ISENTO, P_LIMITE = ("'Parâmetros'!$B$2", "'Parâmetros'!$B$3",
                                         "'Parâmetros'!$B$4", "'Parâmetros'!$B$5")

    # ---- Adições e exclusões (a Parte A soma daqui com SOMASES)
    lanc = wb.create_sheet("Adições e exclusões")
    cabecalho(lanc)
    lanc.append(["Período", "Tipo", "Descrição", "Tributo", "Valor (R$)"])
    for c in lanc[lanc.max_row]:
        c.font, c.fill = negrito, fundo
    for p, d in periodos:
        for l in d.apuracao.lancamentos:
            lanc.append([rotulo_periodo(p, curto=True), LalurLancamento.TIPOS[l.tipo], l.descricao,
                         LalurLancamento.TRIBUTOS[l.tributo], l.valor])
            lanc.cell(lanc.max_row, 5).number_format = numero
    for letra, w in zip("ABCDE", [12, 10, 50, 14, 16]):
        lanc.column_dimensions[letra].width = w
    L = "'Adições e exclusões'!"

    # ---- Parte A com fórmulas
    cabecalho(ws)
    n = len(periodos)
    col_formula = n + 2
    ws.append(["R$"] + [rotulo_periodo(p, curto=True) for p, _ in periodos] + ["Fórmula / critério"])
    for c in ws[ws.max_row]:
        c.font, c.fill = negrito, fundo
        c.alignment = Alignment(horizontal="center")
    linha_titulo = ws.max_row
    linhas = {}  # chave -> nº da linha

    def secao_(titulo):
        ws.append([titulo])
        for c in ws[ws.max_row][:col_formula]:
            c.fill, c.font = secao, negrito

    def linha(chave, rotulo, celulas, explicacao, estilo="", formato=numero, digitavel=False):
        """celulas: função (índice do período, letra da coluna) -> valor ou fórmula."""
        ws.append([rotulo])
        r = ws.max_row
        linhas[chave] = r
        for i in range(n):
            letra = get_column_letter(i + 2)
            c = ws.cell(r, i + 2, celulas(i, letra))
            c.number_format = formato
            if digitavel:
                c.fill = entrada
            if estilo:
                c.font = negrito
        if estilo:
            ws.cell(r, 1).font = negrito
        ws.cell(r, col_formula, explicacao).font = cinza

    def ref(chave, letra):
        return f"{letra}{linhas[chave]}"

    def somases(letra, tipo, excluir):
        rot = f"{letra}${linha_titulo}"
        return (f'=SUMIFS({L}$E:$E,{L}$A:$A,{rot},{L}$B:$B,"{tipo}",{L}$D:$D,"<>{excluir}")')

    ap = [d.apuracao for _, d in periodos]
    secao_("Dados do período")
    linha("lucro", "Lucro líquido antes do IRPJ/CSLL", lambda i, c: ap[i].lucro_antes,
          "Informado (resultado contábil do período)", digitavel=True)
    linha("meses", "Meses do período", lambda i, c: 12 if periodos[i][0] == 0 else 3,
          "3 no trimestre, 12 na apuração anual", formato="0", digitavel=True)
    linha("comp", "Compensar prejuízos (1 = sim, 0 = não)", lambda i, c: 1 if ap[i].compensar else 0,
          "Informado", formato="0", digitavel=True)
    linha("aliq", "Alíquota da CSLL (%)", lambda i, c: ap[i].aliquota_csll, "Informado", digitavel=True)

    secao_("IRPJ")
    linha("i_ad", "(+) Adições", lambda i, c: somases(c, "Adição", "Só CSLL"),
          "Soma das adições do período (IRPJ e CSLL + Só IRPJ) na aba Adições e exclusões")
    linha("i_ex", "(−) Exclusões", lambda i, c: somases(c, "Exclusão", "Só CSLL"),
          "Soma das exclusões do período (IRPJ e CSLL + Só IRPJ)")
    linha("i_aj", "= Lucro ajustado", lambda i, c: f"={ref('lucro', c)}+{ref('i_ad', c)}-{ref('i_ex', c)}",
          "Lucro antes + Adições − Exclusões", "subtotal")
    linha("i_sa", "Prejuízo fiscal a compensar (saldo anterior)",
          lambda i, c: periodos[0][1].irpj.saldo_anterior if i == 0 else f"={get_column_letter(i + 1)}{{i_sf}}",
          "1º período: saldo da Parte B; demais: saldo final do período anterior")
    linha("i_co", "(−) Compensação de prejuízo fiscal",
          lambda i, c: f"=IF(AND({ref('comp', c)}=1,{ref('i_aj', c)}>0),"
                       f"MIN({ref('i_sa', c)},ROUND({ref('i_aj', c)}*{P_LIMITE},2)),0)",
          "Menor entre o saldo e 30% do lucro ajustado (se lucro ajustado positivo)")
    linha("i_lr", "= Lucro real", lambda i, c: f"=MAX({ref('i_aj', c)}-{ref('i_co', c)},0)",
          "Lucro ajustado − Compensação (zero se negativo)", "subtotal")
    linha("i_15", "IRPJ 15%", lambda i, c: f"=ROUND({ref('i_lr', c)}*{P_IRPJ},2)", "Lucro real × 15%")
    linha("i_10", "Adicional 10%",
          lambda i, c: f"=ROUND(MAX({ref('i_lr', c)}-{P_ISENTO}*{ref('meses', c)},0)*{P_ADIC},2)",
          "(Lucro real − R$ 20.000,00 × meses) × 10%, se positivo")
    linha("i_dv", "IRPJ devido", lambda i, c: f"={ref('i_15', c)}+{ref('i_10', c)}", "IRPJ 15% + Adicional",
          "subtotal")
    linha("i_de", "(−) Deduções (retenções/antecipações)", lambda i, c: ap[i].irpj_deduzir, "Informado",
          digitavel=True)
    linha("i_pg", "IRPJ a pagar", lambda i, c: f"=MAX({ref('i_dv', c)}-{ref('i_de', c)},0)",
          "IRPJ devido − Deduções (zero se negativo)", "total")
    linha("i_ge", "Prejuízo fiscal gerado no período", lambda i, c: f"=MAX(-{ref('i_aj', c)},0)",
          "Lucro ajustado negativo vira prejuízo a compensar")
    linha("i_sf", "Prejuízo fiscal a compensar (saldo final)",
          lambda i, c: f"={ref('i_sa', c)}-{ref('i_co', c)}+{ref('i_ge', c)}",
          "Saldo anterior − Compensação + Prejuízo gerado", "subtotal")

    secao_("CSLL")
    linha("c_ad", "(+) Adições", lambda i, c: somases(c, "Adição", "Só IRPJ"),
          "Soma das adições do período (IRPJ e CSLL + Só CSLL)")
    linha("c_ex", "(−) Exclusões", lambda i, c: somases(c, "Exclusão", "Só IRPJ"),
          "Soma das exclusões do período (IRPJ e CSLL + Só CSLL)")
    linha("c_aj", "= Base ajustada", lambda i, c: f"={ref('lucro', c)}+{ref('c_ad', c)}-{ref('c_ex', c)}",
          "Lucro antes + Adições − Exclusões", "subtotal")
    linha("c_sa", "Base negativa a compensar (saldo anterior)",
          lambda i, c: periodos[0][1].csll.saldo_anterior if i == 0 else f"={get_column_letter(i + 1)}{{c_sf}}",
          "1º período: saldo da Parte B; demais: saldo final do período anterior")
    linha("c_co", "(−) Compensação de base negativa",
          lambda i, c: f"=IF(AND({ref('comp', c)}=1,{ref('c_aj', c)}>0),"
                       f"MIN({ref('c_sa', c)},ROUND({ref('c_aj', c)}*{P_LIMITE},2)),0)",
          "Menor entre o saldo e 30% da base ajustada (se positiva)")
    linha("c_bc", "= Base de cálculo", lambda i, c: f"=MAX({ref('c_aj', c)}-{ref('c_co', c)},0)",
          "Base ajustada − Compensação (zero se negativa)", "subtotal")
    linha("c_dv", "CSLL devida", lambda i, c: f"=ROUND({ref('c_bc', c)}*{ref('aliq', c)}/100,2)",
          "Base de cálculo × alíquota da CSLL", "subtotal")
    linha("c_de", "(−) Deduções (retenções/antecipações)", lambda i, c: ap[i].csll_deduzir, "Informado",
          digitavel=True)
    linha("c_pg", "CSLL a pagar", lambda i, c: f"=MAX({ref('c_dv', c)}-{ref('c_de', c)},0)",
          "CSLL devida − Deduções (zero se negativo)", "total")
    linha("c_ge", "Base negativa gerada no período", lambda i, c: f"=MAX(-{ref('c_aj', c)},0)",
          "Base ajustada negativa vira base negativa a compensar")
    linha("c_sf", "Base negativa a compensar (saldo final)",
          lambda i, c: f"={ref('c_sa', c)}-{ref('c_co', c)}+{ref('c_ge', c)}",
          "Saldo anterior − Compensação + Base negativa gerada", "subtotal")

    secao_("Resultado do período")
    linha("ll", "= Lucro líquido do período",
          lambda i, c: f"={ref('lucro', c)}-{ref('i_dv', c)}-{ref('c_dv', c)}",
          "Lucro antes − IRPJ devido − CSLL devida", "total")

    # saldos anteriores dos períodos seguintes apontam para o saldo final da coluna anterior
    for r in (linhas["i_sa"], linhas["c_sa"]):
        for i in range(1, n + 1):
            c = ws.cell(r, i + 1)
            if isinstance(c.value, str):
                c.value = c.value.replace("{i_sf}", str(linhas["i_sf"])).replace("{c_sf}", str(linhas["c_sf"]))

    ws.column_dimensions["A"].width = 46
    for i in range(n):
        ws.column_dimensions[get_column_letter(i + 2)].width = 16
    ws.column_dimensions[get_column_letter(col_formula)].width = 70
    ws.freeze_panes = ws.cell(linha_titulo + 1, 2)
    ws.cell(ws.max_row + 2, 1, "Células em amarelo são dados informados; as demais são fórmulas.").font = cinza

    wb.move_sheet("Parâmetros", offset=1)
    wb.calculation = CalcProperties(fullCalcOnLoad=True)
    saida = BytesIO()
    wb.save(saida)
    return saida.getvalue()


@bp.route("/<int:id>/<int:ano>/exportar")
def exportar(id, ano):
    """?periodo=todos (todos os apurados do exercício) ou 0-4; ?formato=pdf ou xlsx."""
    from flask import Response

    empresa = _empresa(id)
    if (r := _exige_lucro_real(empresa)):
        return r
    calc, _ = demonstrativos(empresa)
    do_ano = sorted(((p, d) for (a, p), d in calc.items() if a == ano), key=lambda x: x[0] or 5)
    escolha = request.args.get("periodo", "todos")
    periodos = do_ano if escolha == "todos" else [(p, d) for p, d in do_ano if str(p) == escolha]
    if not periodos:
        flash("Não há período apurado para exportar (informe o lucro e salve a apuração).", "warning")
        return redirect(url_for("lalur.empresa", id=id, ano=ano))
    if request.args.get("formato") == "xlsx":
        conteudo = _excel(empresa, ano, periodos)
        tipo, extensao = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "xlsx"
    else:
        conteudo = _pdf(empresa, ano, periodos)
        tipo, extensao = "application/pdf", "pdf"
    nome = _nome_arquivo(empresa, ano, periodos, extensao)
    return Response(conteudo, mimetype=tipo,
                    headers={"Content-Disposition": f'attachment; filename="{nome}"'})


# ---------------------------------------------------------------- dados para o DARF

@bp.route("/<int:id>/<int:ano>/<int:periodo>/guia")
def guia(id, ano, periodo):
    """Quota única ou parcelado: quotas, vencimentos, juros e campos do DARF (tela e PDF)."""
    from flask import Response

    from . import guia_darf

    empresa = _empresa(id)
    if (r := _exige_lucro_real(empresa)):
        return r
    periodo = _periodo(periodo)
    calc, _ = demonstrativos(empresa)
    d = calc.get((ano, periodo))
    if d is None or (d.irpj_a_pagar <= 0 and d.csll_a_pagar <= 0):
        flash("Não há IRPJ nem CSLL a pagar neste período.", "warning")
        return _voltar_periodo(id, ano, periodo)
    a = request.args
    anual = periodo == 0
    forma = "unica" if anual or a.get("forma") != "parcelado" else "parcelado"
    try:
        parcelas = 1 if forma == "unica" else int(a.get("parcelas", "3"))
    except ValueError:
        parcelas = 3
    try:
        selic = parse_valor(a.get("selic"))
    except ValueError:
        selic = None
        flash("Selic inválida: informe o percentual, como 1,05.", "warning")
    padrao = guia_darf.CODIGOS["anual" if anual else "trimestral"]
    codigos = {t: (a.get(f"cod_{t}") or padrao[t]).strip()[:6] for t in ("irpj", "csll")}
    guias, avisos = [], []
    for tributo, chave, valor in (("IRPJ", "irpj", d.irpj_a_pagar), ("CSLL", "csll", d.csll_a_pagar)):
        lista, aviso = guia_darf.quotas(valor, ano, periodo, parcelas, selic)
        if lista:
            guias.append((tributo, codigos[chave], lista))
        if aviso:
            avisos.append(f"{tributo}: {aviso}")
    if a.get("baixar"):
        conteudo = guia_darf.pdf(empresa, ano, periodo, guias, forma)
        nome = (f"darf_dados_{empresa.codigo or empresa.id}_{ano}_"
                f"{'anual' if anual else f'{periodo}tri'}_{forma}.pdf")
        return Response(conteudo, mimetype="application/pdf",
                        headers={"Content-Disposition": f'attachment; filename="{nome}"'})
    return render_template(
        "lalur/guia.html", empresa=empresa, ano=ano, periodo=periodo, rotulo=rotulo_periodo(periodo),
        d=d, forma=forma, parcelas=parcelas, selic=selic, codigos=codigos, guias=guias,
        avisos=avisos, anual=anual, apuracao=guia_darf.fim_do_periodo(ano, periodo),
        sicalc=guia_darf.SICALC,
    )


# ---------------------------------------------------------------- lucro do grupo econômico

def _lucro_do_grupo(grupo, ano):
    """Empresas ativas do grupo e o lucro antes do IRPJ/CSLL informado no LALUR por período.

    Devolve (empresas, colunas, valores) com colunas = [1, 2, 3, 4] (+ 0 se houver apuração
    anual no ano) e valores[(empresa_id, periodo)] = Decimal.
    """
    def chave(e):
        c = e.codigo or ""
        return (0, int(c), "") if c.isdigit() else (1, 0, c or e.razao_social)

    empresas = sorted((e for e in grupo.empresas if e.ativo), key=chave)
    ids = [e.id for e in empresas]
    valores = {}
    if ids:
        for a in LalurApuracao.query.filter(LalurApuracao.empresa_id.in_(ids), LalurApuracao.ano == ano,
                                            LalurApuracao.lucro_antes.isnot(None)):
            valores[(a.empresa_id, a.periodo)] = a.lucro_antes
    colunas = [1, 2, 3, 4] + ([0] if any(p == 0 for _, p in valores) else [])
    return empresas, colunas, valores


def _rotulo_coluna_grupo(p):
    return "ANUAL" if p == 0 else f"{p}º TRI"


def _excel_grupo(grupo, ano, empresas, colunas, valores):
    from io import BytesIO

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.workbook.properties import CalcProperties

    wb = Workbook()
    ws = wb.active
    ws.title = f"Grupo {grupo.nome}"[:31]
    numero = "#,##0.00;-#,##0.00"
    ws.append([f"Lucro antes do IRPJ/CSLL — grupo econômico {grupo.nome} — {ano}"])
    ws["A1"].font = Font(bold=True, size=12)
    ws.append([])
    cab = ["CODIGO", "RAZÃO SOCIAL", "CNPJ", "GRUPO"] + [_rotulo_coluna_grupo(p) for p in colunas] + ["TOTAL"]
    ws.append(cab)
    fino = Side(style="thin", color="999999")
    borda = Border(top=fino, bottom=fino, left=fino, right=fino)
    for c in ws[3]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="8B1E2D")
        c.alignment = Alignment(horizontal="center")
        c.border = borda
    primeira = 4
    col_valor1 = 5
    col_total = col_valor1 + len(colunas)
    for e in empresas:
        codigo = e.codigo or ""
        ws.append([int(codigo) if codigo.isdigit() else codigo, e.razao_social, e.cnpj_formatado, grupo.nome]
                  + [valores.get((e.id, p)) for p in colunas])
        r = ws.max_row
        ini, fim = get_column_letter(col_valor1), get_column_letter(col_total - 1)
        ws.cell(r, col_total, f"=SUM({ini}{r}:{fim}{r})")
        for c in ws[r]:
            c.border = borda
        for i in range(col_valor1, col_total + 1):
            ws.cell(r, i).number_format = numero
            if i < col_total:
                ws.cell(r, i).fill = PatternFill("solid", fgColor="DDEBF7")
    ultima = ws.max_row
    ws.append(["", "TOTAL DO GRUPO"])
    r = ws.max_row
    for i in range(col_valor1, col_total + 1):
        letra = get_column_letter(i)
        ws.cell(r, i, f"=SUM({letra}{primeira}:{letra}{ultima})" if empresas else 0).number_format = numero
    for c in ws[r]:
        c.font = Font(bold=True)
        c.border = borda
    ws.append([])
    ws.append(["Valores: lucro líquido antes do IRPJ e da CSLL informado no LALUR de cada período; "
               "em branco = período não apurado."])
    ws.cell(ws.max_row, 1).font = Font(italic=True, color="666666")
    for letra, w in zip("ABCD", [10, 52, 20, 14]):
        ws.column_dimensions[letra].width = w
    for i in range(col_valor1, col_total + 1):
        ws.column_dimensions[get_column_letter(i)].width = 16
    ws.freeze_panes = "C4"
    wb.calculation = CalcProperties(fullCalcOnLoad=True)
    saida = BytesIO()
    wb.save(saida)
    return saida.getvalue()


def _pdf_grupo(grupo, ano, empresas, colunas, valores):
    from .competencia import fmt_valor
    from .pdf_simples import A4, Documento

    doc = Documento(f"Lucro do grupo {grupo.nome} {ano}", tamanho=(A4[1], A4[0]))  # paisagem
    m = 30
    larg = doc.largura_pagina - 2 * m
    fixas = [46, 0, 104, 60]
    valor_w = 78
    n_val = len(colunas) + 1
    fixas[1] = larg - sum(fixas) - valor_w * n_val
    larguras = fixas + [valor_w] * n_val
    cab = ["CÓDIGO", "RAZÃO SOCIAL", "CNPJ", "GRUPO"] + [_rotulo_coluna_grupo(p) for p in colunas] + ["TOTAL"]
    alt = 16
    pagina = [0]

    def topo():
        pagina[0] += 1
        doc.texto(m, 34, f"Lucro antes do IRPJ/CSLL — grupo econômico {grupo.nome}", 13, negrito=True)
        doc.texto(m + larg, 34, f"Exercício {ano}", 11, negrito=True, alinhar="direita")
        doc.texto(m, 48, "Lucro líquido antes do IRPJ e da CSLL informado no LALUR de cada período "
                         "(em branco = não apurado).", 8.5, cinza=0.35)
        doc.texto(m, doc.altura_pagina - 18, f"Emitido em {date.today():%d/%m/%Y} pelo Controle de Rotina",
                  7.5, cinza=0.45)
        doc.texto(m + larg, doc.altura_pagina - 18, f"Página {pagina[0]}", 7.5, alinhar="direita", cinza=0.45)
        y = 60
        doc.faixa(m, y, larg, alt, 0.25)
        x = m
        for i, (c, w) in enumerate(zip(cab, larguras)):
            if i < 4:
                doc.texto(x + 4, y + 11, c, 8, negrito=True, cinza=1)
            else:
                doc.texto(x + w - 4, y + 11, c, 8, negrito=True, cinza=1, alinhar="direita")
            x += w
        return y + alt

    def linha(y, celulas, negrito=False, destaque=False):
        if destaque:
            doc.faixa(m, y, larg, alt, 0.92)
        x = m
        for i, (c, w) in enumerate(zip(celulas, larguras)):
            if i < 4:
                texto = c
                while texto and largura_texto(texto, 8.5, negrito) > w - 8:
                    texto = texto[:-2] + "…"
                doc.texto(x + 4, y + 11.5, texto, 8.5, negrito=negrito)
            else:
                if i - 4 < len(colunas) and not negrito:
                    doc.faixa(x, y, w, alt, 0.93)
                doc.texto(x + w - 4, y + 11.5, c, 8.5, negrito=negrito, alinhar="direita")
            x += w
        doc.linha(m, y + alt, m + larg, y + alt, 0.3, 0.75)
        return y + alt

    from .pdf_simples import largura as largura_texto

    y = topo()
    totais = [Decimal("0")] * n_val
    for e in empresas:
        if y + alt > doc.altura_pagina - 30:
            doc.nova_pagina()
            y = topo()
        vals = [valores.get((e.id, p)) for p in colunas]
        total = sum((v for v in vals if v is not None), Decimal("0"))
        for i, v in enumerate(vals + [total]):
            totais[i] += v or 0
        y = linha(y, [e.codigo or "", e.razao_social, e.cnpj_formatado, grupo.nome]
                  + [fmt_valor(v) for v in vals] + [fmt_valor(total)])
    if y + alt > doc.altura_pagina - 30:
        doc.nova_pagina()
        y = topo()
    linha(y, ["", "TOTAL DO GRUPO", "", ""] + [fmt_valor(t) for t in totais], negrito=True, destaque=True)
    return doc.gerar()


@bp.route("/grupo")
def grupo():
    """Escolha do grupo econômico e do ano; prévia e exportação em Excel/PDF."""
    from flask import Response

    from .models import GrupoEconomico

    grupos = GrupoEconomico.query.order_by(GrupoEconomico.nome).all()
    escolhido = db.session.get(GrupoEconomico, request.args.get("grupo", type=int) or 0)
    ano = _ano()
    empresas, colunas, valores = _lucro_do_grupo(escolhido, ano) if escolhido else ([], [1, 2, 3, 4], {})
    formato = request.args.get("formato")
    if escolhido and formato in ("xlsx", "pdf"):
        if formato == "xlsx":
            conteudo = _excel_grupo(escolhido, ano, empresas, colunas, valores)
            tipo = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        else:
            conteudo = _pdf_grupo(escolhido, ano, empresas, colunas, valores)
            tipo = "application/pdf"
        nome_grupo = "".join(ch if ch.isalnum() else "_" for ch in escolhido.nome)[:30]
        return Response(conteudo, mimetype=tipo, headers={
            "Content-Disposition": f'attachment; filename="lucro_grupo_{nome_grupo}_{ano}.{formato}"'})
    linhas = []
    for e in empresas:
        vals = [valores.get((e.id, p)) for p in colunas]
        linhas.append((e, vals, sum((v for v in vals if v is not None), Decimal("0"))))
    totais = [sum((l[1][i] or 0 for l in linhas), Decimal("0")) for i in range(len(colunas))]
    return render_template("lalur/grupo.html", grupos=grupos, escolhido=escolhido, ano=ano,
                           colunas=[(p, _rotulo_coluna_grupo(p)) for p in colunas], linhas=linhas,
                           totais=totais, total_geral=sum(totais, Decimal("0")))
