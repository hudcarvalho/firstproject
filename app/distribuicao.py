"""Aba Distribuição de Lucros: sócios de cada cliente e valores distribuídos por mês."""
from collections import defaultdict
from datetime import date
from decimal import Decimal

from flask import Blueprint, abort, flash, redirect, render_template, request, url_for
from sqlalchemy import or_

from .competencia import cnpj_valido, cpf_valido, fmt_documento, fmt_valor, parse_valor, so_digitos
from .models import DistribuicaoLucro, Empresa, Socio, db

bp = Blueprint("distribuicao", __name__, url_prefix="/distribuicao")

MESES = ["Jan", "Fev", "Mar", "Abr", "Mai", "Jun", "Jul", "Ago", "Set", "Out", "Nov", "Dez"]


def _empresa(id_):
    empresa = db.session.get(Empresa, id_)
    if empresa is None:
        abort(404)
    return empresa


def _ano():
    try:
        ano = int(request.values.get("ano", ""))
    except ValueError:
        return date.today().year
    return ano if 2000 <= ano <= 2100 else date.today().year


def _percentual(texto):
    """Percentual digitado ("33,33", "50", "12.5") -> Decimal entre 0 e 100, ou None se vazio.

    Levanta ValueError se inválido.
    """
    valor = parse_valor((texto or "").replace("%", ""))
    if valor is not None and not Decimal("0") <= valor <= Decimal("100"):
        raise ValueError(texto)
    return valor


def _soma_quotas(empresa):
    return sum((s.percentual or Decimal("0") for s in empresa.socios if s.ativo), Decimal("0"))


def _avisar_soma(empresa):
    ativos = [s for s in empresa.socios if s.ativo]
    if ativos and all(s.percentual is not None for s in ativos):
        soma = _soma_quotas(empresa)
        if soma != Decimal("100"):
            flash(f"Atenção: as quotas dos sócios ativos somam {fmt_valor(soma)}% (e não 100%).",
                  "warning")


def _clientes_para_busca():
    return Empresa.query.filter_by(ativo=True).order_by(Empresa.razao_social).all()


@bp.route("/")
def selecionar():
    """Escolha do cliente. Com ?busca=, resolve código exato ou nome e vai direto à empresa."""
    busca = (request.args.get("busca") or "").strip()
    encontrados = []
    if busca:
        rotulo = busca.split(" — ", 1)  # valor vindo da lista: "código — razão social"
        codigo = rotulo[0].strip()
        por_codigo = Empresa.query.filter_by(codigo=codigo).all() if codigo else []
        if len(rotulo) == 2:
            por_codigo = [e for e in por_codigo if e.razao_social == rotulo[1].strip()] or por_codigo
        if len(por_codigo) == 1:
            return redirect(url_for("distribuicao.empresa", id=por_codigo[0].id))
        digitos = so_digitos(busca)
        filtros = [Empresa.razao_social.ilike(f"%{busca}%"), Empresa.nome_fantasia.ilike(f"%{busca}%")]
        if len(digitos) >= 5:
            filtros.append(Empresa.cnpj.contains(digitos))
        encontrados = por_codigo or Empresa.query.filter(or_(*filtros)).order_by(
            Empresa.razao_social).limit(50).all()
        if len(encontrados) == 1:
            return redirect(url_for("distribuicao.empresa", id=encontrados[0].id))
        if not encontrados:
            flash(f"Nenhum cliente encontrado para “{busca}”.", "warning")
    return render_template("distribuicao/selecionar.html", clientes=_clientes_para_busca(),
                           busca=busca, encontrados=encontrados)


def _grade(empresa, ano):
    """Linhas da grade: sócios ativos + inativos que tenham valores no ano."""
    valores = defaultdict(dict)  # socio_id -> {mes: valor}
    ids = [s.id for s in empresa.socios]
    if ids:
        for d in DistribuicaoLucro.query.filter(DistribuicaoLucro.socio_id.in_(ids),
                                                DistribuicaoLucro.ano == ano):
            valores[d.socio_id][d.mes] = d.valor
    linhas = []
    for s in empresa.socios:
        if s.ativo or valores.get(s.id):
            meses = [valores[s.id].get(m) for m in range(1, 13)]
            linhas.append({"socio": s, "meses": meses,
                           "total": sum((v for v in meses if v is not None), Decimal("0"))})
    totais_mes = [sum((l["meses"][i] or Decimal("0") for l in linhas), Decimal("0")) for i in range(12)]
    return linhas, totais_mes, sum(totais_mes, Decimal("0"))


@bp.route("/<int:id>", methods=["GET", "POST"])
def empresa(id):
    empresa = _empresa(id)
    ano = _ano()
    if request.method == "POST":
        erros = _salvar_valores(empresa, ano)
        if erros:
            db.session.rollback()
            for e in erros:
                flash(e, "danger")
        else:
            db.session.commit()
            flash(f"Valores de {ano} salvos.", "success")
            return redirect(url_for("distribuicao.empresa", id=id, ano=ano))
    linhas, totais_mes, total_geral = _grade(empresa, ano)
    return render_template(
        "distribuicao/empresa.html", empresa=empresa, ano=ano, meses=MESES, linhas=linhas,
        totais_mes=totais_mes, total_geral=total_geral,
        soma_quotas=_soma_quotas(empresa),
        valores_digitados=request.form if request.method == "POST" else None,
    )


def _salvar_valores(empresa, ano):
    """Campos do formulário: v-<socio_id>-<mes>. Vazio apaga o valor do mês."""
    socios = {s.id: s for s in empresa.socios}
    existentes = {}
    if socios:
        for d in DistribuicaoLucro.query.filter(DistribuicaoLucro.socio_id.in_(socios),
                                                DistribuicaoLucro.ano == ano):
            existentes[(d.socio_id, d.mes)] = d
    erros = []
    for chave, texto in request.form.items():
        partes = chave.split("-")
        if len(partes) != 3 or partes[0] != "v":
            continue
        try:
            socio_id, mes = int(partes[1]), int(partes[2])
        except ValueError:
            continue
        if socio_id not in socios or not 1 <= mes <= 12:
            continue
        try:
            valor = parse_valor(texto)
        except ValueError:
            erros.append(f"{socios[socio_id].nome} — {MESES[mes - 1]}: “{texto}” não é um valor válido.")
            continue
        if valor is not None and valor < 0:
            erros.append(f"{socios[socio_id].nome} — {MESES[mes - 1]}: o valor não pode ser negativo.")
            continue
        atual = existentes.get((socio_id, mes))
        if valor is None or valor == 0:
            if atual:
                db.session.delete(atual)
        elif atual:
            atual.valor = valor
        else:
            db.session.add(DistribuicaoLucro(socio_id=socio_id, ano=ano, mes=mes, valor=valor))
    return erros


@bp.post("/<int:id>/socios")
def socio_novo(id):
    empresa = _empresa(id)
    ano = _ano()
    nome = " ".join((request.form.get("nome") or "").split())
    documento = so_digitos(request.form.get("documento"))
    try:
        percentual = _percentual(request.form.get("percentual"))
    except ValueError:
        flash("Quotas: informe um percentual entre 0 e 100.", "danger")
        return redirect(url_for("distribuicao.empresa", id=id, ano=ano))
    if not nome:
        flash("Informe o nome do sócio.", "danger")
    elif not (cpf_valido(documento) or cnpj_valido(documento)):
        flash("CPF/CNPJ inválido.", "danger")
    elif Socio.query.filter_by(empresa_id=id, documento=documento).first():
        flash("Já existe um sócio com esse CPF/CNPJ nesta empresa.", "warning")
    else:
        db.session.add(Socio(empresa=empresa, nome=nome, documento=documento,
                             percentual=percentual, ativo=True))
        db.session.commit()
        flash(f"Sócio {nome} cadastrado.", "success")
        _avisar_soma(empresa)
    return redirect(url_for("distribuicao.empresa", id=id, ano=ano))


@bp.post("/<int:id>/quotas")
def quotas(id):
    """Grava o % de quotas de cada sócio (campos pct-<socio_id>)."""
    empresa = _empresa(id)
    socios = {s.id: s for s in empresa.socios}
    novos, erros = {}, []
    for chave, texto in request.form.items():
        if not chave.startswith("pct-"):
            continue
        try:
            socio = socios[int(chave[4:])]
        except (ValueError, KeyError):
            continue
        try:
            novos[socio] = _percentual(texto)
        except ValueError:
            erros.append(f"{socio.nome}: “{texto}” não é um percentual entre 0 e 100.")
    if erros:
        for e in erros:
            flash(e, "danger")
    else:
        for socio, valor in novos.items():
            socio.percentual = valor
        db.session.commit()
        flash("Quotas salvas.", "success")
        _avisar_soma(empresa)
    return redirect(url_for("distribuicao.empresa", id=id, ano=_ano()))


@bp.post("/socios/<int:socio_id>/situacao")
def socio_situacao(socio_id):
    socio = db.session.get(Socio, socio_id)
    if socio is None:
        abort(404)
    socio.ativo = not socio.ativo
    db.session.commit()
    flash(f"Sócio {socio.nome} {'reativado' if socio.ativo else 'inativado'}.", "success")
    return redirect(url_for("distribuicao.empresa", id=socio.empresa_id, ano=_ano()))


# ---------------------------------------------------------------- Exportação para Excel

COLUNAS_EXPORTACAO = ["CODIGO", "RAZÃO SOCIAL", "CNPJ", "Nome do sócio", "Sócio CPF/CNPJ",
                      "Distribuição", "REF."]


def _mes(nome, padrao):
    try:
        mes = int(request.args.get(nome, ""))
    except ValueError:
        return padrao
    return mes if 1 <= mes <= 12 else padrao


def _lancamentos(ano, mes_de, mes_ate, empresa_id=None):
    """Distribuições do período (valores diferentes de zero), por cliente, sócio e mês."""
    consulta = (
        db.session.query(DistribuicaoLucro, Socio, Empresa)
        .join(Socio, DistribuicaoLucro.socio_id == Socio.id)
        .join(Empresa, Socio.empresa_id == Empresa.id)
        .filter(DistribuicaoLucro.ano == ano, DistribuicaoLucro.mes.between(mes_de, mes_ate),
                DistribuicaoLucro.valor != 0)
    )
    if empresa_id:
        consulta = consulta.filter(Empresa.id == empresa_id)
    return consulta.order_by(Empresa.razao_social, Empresa.id, Socio.nome, Socio.id,
                             DistribuicaoLucro.mes).all()


def _planilha(lancamentos):
    from io import BytesIO

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    wb = Workbook()
    ws = wb.active
    ws.title = "Distribuição de lucros"
    ws.append(COLUNAS_EXPORTACAO)
    fino = Side(style="thin")
    for cel in ws[1]:
        cel.font = Font(bold=True)
        cel.alignment = Alignment(horizontal="center", vertical="center")
        cel.border = Border(top=fino, bottom=fino, left=fino, right=fino)
        cel.fill = PatternFill("solid", fgColor="D9E1F2")
    for d, socio, empresa in lancamentos:
        codigo = empresa.codigo or ""
        ws.append([
            int(codigo) if codigo.isdigit() else codigo,
            empresa.razao_social,
            empresa.cnpj_formatado,
            socio.nome,
            fmt_documento(socio.documento),
            d.valor,
            f"{d.mes:02d}/{d.ano}",
        ])
        ws.cell(ws.max_row, 6).number_format = "#,##0.00"
    for letra, largura in zip("ABCDEFG", [10, 50, 20, 40, 20, 16, 10]):
        ws.column_dimensions[letra].width = largura
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:G{ws.max_row}"
    saida = BytesIO()
    wb.save(saida)
    saida.seek(0)
    return saida


@bp.route("/exportar")
def exportar():
    """Tela para escolher o período e baixar as distribuições em Excel."""
    from flask import send_file

    ano = _ano()
    mes_de = _mes("de", 1)
    mes_ate = max(_mes("ate", 12), mes_de)
    empresa = db.session.get(Empresa, request.args.get("empresa", type=int) or 0)
    lancamentos = _lancamentos(ano, mes_de, mes_ate, empresa.id if empresa else None)
    if request.args.get("baixar"):
        if not lancamentos:
            flash("Nenhuma distribuição lançada nesse período.", "warning")
        else:
            nome = f"distribuicao_lucros_{ano}"
            if (mes_de, mes_ate) != (1, 12):
                nome += f"_{mes_de:02d}" + (f"-{mes_ate:02d}" if mes_ate != mes_de else "")
            if empresa:
                nome += f"_{empresa.codigo or empresa.id}"
            return send_file(
                _planilha(lancamentos), as_attachment=True, download_name=nome + ".xlsx",
                mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
    anos = sorted({a for (a,) in db.session.query(DistribuicaoLucro.ano).distinct()}
                  | {date.today().year}, reverse=True)
    return render_template(
        "distribuicao/exportar.html", ano=ano, anos=anos, mes_de=mes_de, mes_ate=mes_ate,
        meses=MESES, empresa=empresa, quantidade=len(lancamentos),
        clientes=len({e.id for _, _, e in lancamentos}),
        total=sum((d.valor for d, _, _ in lancamentos), Decimal("0")),
    )
